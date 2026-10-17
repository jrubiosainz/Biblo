"""More of Biblo's demo files, the ones to read: seven PDFs, four Word documents, an RTF glossary and four decks.

make_demo.py imports this module, so its builders are groups like any other (scripts/make-demo.sh more_documents).
They belong to the same made-up story as more_sound.py and more_pictures.py: the PDFs the returns assistant cites,
Trailhead 3's launch, the offsite in Lisbon, the Data Summit in Madrid and GitHub Copilot's rollout. Every fact about a
Microsoft or GitHub product is public; every company, person and number is invented. Logos appear only on slides, whole
and unaltered (assets/SOURCES.md)."""
from __future__ import annotations

import datetime as dt
import re
from pathlib import Path

from pptx.enum.text import MSO_ANCHOR, PP_ALIGN

from make_demo import (DATES, LEFT, MONO, RIGHT, WORK, Deck, box, builder, dot, esc, link, mark, out, page, picture,
                       place, print_pdf, render, save_png, text, textutil, write)
from more_pictures import dome

DATES.update({
    "Documents/Return policy.pdf": "2026-10-01 10:12",
    "Documents/Shipping options.pdf": "2026-10-13 17:25",
    "Documents/Trailhead 3 – spec sheet.pdf": "2026-10-06 16:30",
    "Documents/Política de viajes y gastos.pdf": "2026-09-24 11:05",
    "Documents/Data Summit Madrid 2026 – agenda.pdf": "2026-10-09 13:05",
    "Documents/KQL cheat sheet.pdf": "2026-10-14 09:40",
    "Documents/Chuleta de prompts para Copilot.pdf": "2026-09-29 10:20",
    "Documents/Nota de prensa – Trailhead 3 (borrador v2).docx": "2026-10-14 18:10",
    "Documents/Agenda offsite Lisboa.docx": "2026-10-13 19:40",
    "Documents/Guion keynote – Data Summit Madrid.docx": "2026-10-15 09:20",
    "Documents/Acta – comité de seguridad (16 oct).docx": "2026-10-16 13:20",
    "Documents/Glosario de IA generativa.rtf": "2026-10-05 11:30",
    "Slides/Data Summit Madrid – keynote.pptx": "2026-10-13 11:40",
    "Slides/Trailhead 3 – launch plan.pptx": "2026-10-15 16:05",
    "Slides/Offsite Lisboa – agenda.pptx": "2026-10-15 13:10",
    "Slides/GitHub Copilot rollout – engineering.pptx": "2026-10-12 17:20",
})

# MARK: PDFs

# A4 sheets, one per page. A sheet whose content doesn't fit grows onto the next page, so a PDF with more pages than
# sheets has one that overflows.
SHEET_CSS = """@page{size:A4;margin:0}
html,body{width:auto;height:auto;overflow:visible}
*{-webkit-print-color-adjust:exact;print-color-adjust:exact}
.sheet{position:relative;min-height:296.5mm;padding:16mm 17mm 22mm;break-after:page}
.sheet:last-child{break-after:auto}
.foot{position:absolute;left:17mm;right:17mm;bottom:10mm;display:flex;justify-content:space-between;font-size:7.5pt}
"""


def sheets(pages: list[str], css: str, title: str, foot: str, path: Path, bg: str = "#fff") -> None:
    """A PDF of A4 sheets, each with `foot` at its bottom and, when there are several, its number."""
    count = len(pages)
    pages = [content.replace(" €", "\u00a0€") for content in pages]
    body = "".join(f'<section class="sheet">{content}<div class="foot"><span>{foot}</span>'
                   f'<span>{f"{number} / {count}" if count > 1 else ""}</span></div></section>'
                   for number, content in enumerate(pages, 1))
    document = page(body, SHEET_CSS + css, bg=bg).replace("<head>", f"<head><title>{esc(title)}</title>", 1)
    print_pdf(document, path)


# Contoso Outdoor's customer care documents: the ones the returns assistant cites, page by page.
TENT = ('<svg viewBox="0 0 28 20" width="30" height="21"><path d="M2 19 L14 2 L26 19 Z" fill="#2f5d46"/>'
        '<path d="M11 19 L14 13 L17 19 Z" fill="#fff"/></svg>')
OUTDOOR_CSS = """body{font-family:"Avenir Next",system-ui;color:#16261f;font-size:10pt;line-height:1.45}
.brand{display:flex;align-items:center;gap:9px;padding-bottom:6mm;margin-bottom:7mm;border-bottom:1.5px solid #2f5d46}
.brand b{font-weight:700;font-size:10pt;letter-spacing:.22em}
.brand span{margin-left:auto;font-size:8.5pt;color:#5b6b62}
h1{margin:0 0 3mm;font-size:30pt;line-height:1.05;letter-spacing:-.01em}
.lede{margin:0 0 7mm;font-size:12.5pt;line-height:1.4;color:#5b6b62;max-width:150mm}
h2{margin:6.5mm 0 2.5mm;font-size:13.5pt;color:#2f5d46}
h2.first{margin-top:0}
h3{margin:0 0 1.5mm;font-size:10.5pt}
p{margin:0 0 2.6mm}
ol,ul{margin:0 0 2.6mm;padding-left:5mm}li{margin-bottom:1.6mm}
table{width:100%;border-collapse:collapse;margin:1mm 0 3mm}
th{text-align:left;font-size:7.5pt;font-weight:700;letter-spacing:.1em;text-transform:uppercase;color:#5b6b62;
padding:2mm 3mm;border-bottom:1.5px solid #16261f}
td{padding:2.6mm 3mm;border-bottom:.6px solid #d5dfd6;vertical-align:top}
.hero{display:flex;gap:8mm;align-items:center;padding:5.5mm 8mm;background:#eef3ee;border-radius:3mm;margin-bottom:4mm}
.hero strong{font-size:46pt;line-height:1;font-weight:700;color:#2f5d46;white-space:nowrap}
.hero p{margin:0;font-size:11.5pt}
.steps{display:grid;grid-template-columns:repeat(4,1fr);gap:4mm}
.steps div{padding:4.5mm 4mm;border:.8px solid #d5dfd6;border-radius:2.5mm}
.steps i{display:grid;place-items:center;width:7mm;height:7mm;border-radius:50%;background:#2f5d46;color:#fff;
font-style:normal;font-weight:700;font-size:9pt;margin-bottom:3mm}
.steps p{margin:0;font-size:9.5pt;color:#3b4a42}
.note{padding:4.5mm 5.5mm;margin:1mm 0 3mm;background:#eef3ee;border-left:1.3mm solid #2f5d46;border-radius:1mm}
.note p:last-child{margin:0}
.fiction{margin-top:5mm;font-size:8pt;color:#5b6b62}
.foot{color:#5b6b62}
.grid{width:100%;border-collapse:separate;border-spacing:1.4mm;margin:0 0 3mm;table-layout:fixed}
.grid th{border:0;padding:0 0 1mm;text-align:center;font-size:8pt;letter-spacing:.06em}
.grid th:first-child{text-align:left}
.grid td{border:0;padding:2mm 0;text-align:center;border-radius:1.5mm;font-size:8.5pt;font-weight:600;
background:#dfeadf;color:#2f5d46}
.grid td.when{background:none;text-align:left;color:#16261f;font-weight:600}
.grid td.off{background:#f3f4f3;color:#a3aca6;font-weight:500}
.top{display:grid;grid-template-columns:1fr 70mm;gap:6mm;align-items:center}
.kicker{margin:0 0 2mm;font-size:8.5pt;font-weight:700;letter-spacing:.18em;text-transform:uppercase;color:#2f5d46}
.art svg{width:70mm;height:auto;display:block}
.spec td{padding:1.6mm 3mm}.spec td:first-child{width:45mm;color:#5b6b62}
.top .lede{margin-bottom:3mm}
.mend{margin:3.5mm 0 0;padding:.4mm 0 .4mm 3.5mm;border-left:1.1mm solid #2f5d46}
.two{display:grid;grid-template-columns:1fr 1fr;gap:9mm}
.swatches{display:flex;gap:5mm}
.swatches div{flex:1;padding:3.5mm;border:.8px solid #d5dfd6;border-radius:2.5mm}
.swatches i{display:block;height:10mm;border-radius:1.5mm;margin-bottom:2.5mm}
.five{display:grid;grid-template-columns:repeat(5,1fr);gap:3mm;padding:0;margin:0 0 4mm;list-style:none;
counter-reset:step}
.five li{margin:0;padding:3.5mm;border:.8px solid #d5dfd6;border-radius:2.5mm;font-size:9pt;line-height:1.35}
.five b{display:block;margin-bottom:1mm;font-size:8pt;letter-spacing:.12em;text-transform:uppercase;color:#2f5d46}
"""


def brand(section: str) -> str:
    return f'<div class="brand">{TENT}<b>CONTOSO OUTDOOR</b><span>{esc(section)}</span></div>'


FICTION = ('<p class="fiction">Contoso Outdoor is a fictitious company, and this document is sample content: its '
           'prices and terms are made up.</p>')


def return_policy() -> None:
    first = f"""{brand("Customer care · Return policy")}
<h1>Return policy</h1>
<p class="lede">Changed your mind, got the wrong size, or something arrived broken? Here's how returns work at Contoso
Outdoor. Updated October 1, 2026.</p>
<div class="hero"><strong>30 days</strong><p>You have <b>30 days from delivery</b> to return anything you bought from
us, online or in a store, for a refund or an exchange. Items must be unused, unless they arrived damaged or are
faulty.</p>
</div>
<h2>What a return costs</h2>
<table><tr><th>Why it comes back</th><th>In a store</th><th>Drop-off point</th><th>Courier pickup</th></tr>
<tr><td><b>Damaged in transit</b>, faulty, or not what you ordered</td><td>Free</td><td>Free</td><td>Free</td></tr>
<tr><td><b>Changed your mind</b> or wrong size</td><td>Free</td><td>4.95 €</td><td>4.95 €</td></tr></table>
<h2>How to return something</h2>
<div class="steps"><div><i>1</i><h3>Tell us</h3><p>Chat with us at contoso.com/help, or start a return at
contoso.com/returns with your order number.</p></div>
<div><i>2</i><h3>Get your RMA</h3><p>We open a return, with an RMA number like RMA-20931, and email you a prepaid
label.</p></div>
<div><i>3</i><h3>Pack it</h3><p>Use the original box if you still have it, and stick the label on the outside.</p></div>
<div><i>4</i><h3>Send it back</h3><p>Take it to a store or a drop-off point, or book a courier pickup from your door:
see Shipping options, page 2.</p></div></div>
<h2>What we can't take back</h2>
<p>Gas canisters and fuel, opened food, worn socks and underwear, and gear that has been altered or cut to size.
Gift cards can't be returned either.</p>"""
    second = f"""{brand("Customer care · Return policy")}
<h2 class="first">Damaged or faulty items</h2>
<p>If something arrives damaged, or stops working because of a defect, returning it is <b>always free</b> within the
30 days, whichever way you send it back.</p>
<div class="note"><p><b>What happens next.</b> As soon as we've found your order, we open a return (an RMA) and email
you a <b>prepaid label</b>. Then you choose: a <b>replacement</b>, sent by express delivery at no cost, or a
<b>refund to your card</b> within 5 working days. Keep the damaged item until the label arrives.</p></div>
<p>Photos help: if you can, send us one of the damage with your request. They're never required.</p>
<h2>Warranty</h2>
<p>Everything we sell has a 2-year warranty against manufacturing defects, and tent poles are covered for 5 years.
Every Contoso Outdoor tent comes with a pole repair sleeve: if a pole breaks on a trip, slide the sleeve over the
break, tape it, and contact us for a new pole.</p>
<h2>Refunds</h2>
<p>We refund the item's price, plus its delivery cost if the whole order comes back, within 5 working days of
receiving the return, to the card you paid with. In stores, card payments go back to the card and cash is refunded in
cash.</p>
<h2>Exchanges</h2>
<p>Want another size or color? Exchanges are free in any store. Online, start a return and place a new order: we
refund the first one as soon as it's back with us.</p>
<h2>Need help?</h2>
<p>Chat with us at contoso.com/help, every day from 8:00 to 22:00, or write to support@contoso.com.</p>{FICTION}"""
    sheets([first, second], OUTDOOR_CSS, "Return policy", "Contoso Outdoor · Return policy · October 2026",
           out("Documents/Return policy.pdf"))


DAYS = ["Mon", "Tue", "Wed", "Thu", "Fri", "Sat", "Sun"]


def shipping_options() -> None:
    first = f"""{brand("Customer care · Shipping options")}
<h1>Shipping options</h1>
<p class="lede">Every order leaves our warehouse in Zaragoza. Order before 14:00, Monday to Friday, and it ships the
same day. Updated October 13, 2026.</p>
<table><tr><th>Option</th><th>Delivery</th><th>Cost</th></tr>
<tr><td><b>Standard</b></td><td>2–3 working days</td><td>3.95 €, free on orders over 60 €</td></tr>
<tr><td><b>Express</b></td><td>1–2 working days, if you order before 14:00</td><td>6.95 €</td></tr>
<tr><td><b>Click &amp; collect</b></td><td>Ready in 2 hours at any store</td><td>Free</td></tr>
<tr><td><b>Parcel locker</b></td><td>2–3 working days, to the locker you choose</td><td>2.95 €</td></tr></table>
<h2>Replacements for damaged items</h2>
<div class="note"><p>When something arrives damaged, its <b>replacement always ships by express delivery, free of
charge</b>, the same day if we've opened your return before 14:00. <b>Keep the damaged item until the prepaid label
arrives</b>: you'll need it to send the item back.</p></div>
<h2>Tracking</h2>
<p>As soon as your order ships, we email you a tracking link. You can also ask our assistant at contoso.com/help:
“Where's my order?” and your order number is enough.</p>
<h2>Balearic and Canary Islands, Ceuta and Melilla</h2>
<p>4–6 working days, 9.95 €. Express isn't available. Orders to the Canary Islands, Ceuta and Melilla may need a
customs declaration, which we prepare for you.</p>
<h2>Public holidays</h2>
<p>We don't ship or deliver on national public holidays, and our delivery estimates already count them.</p>"""
    slots = ["8:00–10:00", "10:00–12:00", "12:00–14:00", "14:00–16:00", "16:00–18:00", "18:00–20:00"]
    header = "".join(f"<th>{day}</th>" for day in DAYS)
    rows = "".join(f'<tr><td class="when">{slot}</td>' + "<td>Pickup</td>" * 6 + '<td class="off">—</td></tr>'
                   for slot in slots)
    second = f"""{brand("Customer care · Shipping options")}
<h2 class="first">Courier pickups for returns</h2>
<p class="lede">No need to go anywhere: a courier can pick up your return at your door. Returns of damaged items get a
free courier pickup.</p>
<div class="hero"><strong>2 h</strong><p>Pickups run <b>Monday to Saturday, from 8:00 to 20:00, in 2-hour slots</b>.
You choose the day and the slot, and we text you 30 minutes before the courier arrives.</p></div>
<table class="grid"><tr><th style="width:24mm">Slot</th>{header}</tr>{rows}</table>
<h2>How to book a pickup</h2>
<ol><li>Chat with us at contoso.com/help and give us your order number. The assistant finds your return and shows you
the free slots.</li>
<li>Book by 20:00 the day before, at the latest.</li>
<li>Have the box closed, with the RMA label on the outside.</li>
<li>We text you 30 minutes before the courier arrives.</li></ol>
<h2>What it costs</h2>
<p>Free for damaged or faulty items, and for items that aren't what you ordered. For any other return it costs
4.95 €, deducted from your refund.</p>
<h2>Changing or cancelling a pickup</h2>
<p>Reply to the confirmation text or ask the assistant, until 20:00 the day before. If the courier misses you, they
come back the next working day in the same slot.</p>{FICTION}"""
    sheets([first, second], OUTDOOR_CSS, "Shipping options", "Contoso Outdoor · Shipping options · October 2026",
           out("Documents/Shipping options.pdf"))


def spec_sheet() -> None:
    specs = [("Sleeps", "3 people"),
             ("Weight", "1.9 kg, 300 g lighter than Trailhead 2"),
             ("Packed size", "46 × 15 cm"),
             ("Floor", "215 × 180 cm · peak height 110 cm"),
             ("Doors and vestibules", "2 doors, 2 vestibules"),
             ("Poles", "2 aluminum poles, color-coded green and sand"),
             ("Fly", "Recycled polyester ripstop, 3,000 mm waterproof"),
             ("Setup", "About 3 minutes, with color-coded poles and clips"),
             ("Warranty", "2 years · poles 5 years"),
             ("Price", "249\u00a0€")]
    table = "".join(f"<tr><td>{name}</td><td><b>{esc(value)}</b></td></tr>" for name, value in specs)
    steps = [("Step 1", "Lay out the inner tent, door away from the wind, and stake the four corners."),
             ("Step 2", "Assemble the poles: green in the green corners, sand in the sand corners."),
             ("Step 3", "Clip the inner tent to the poles, from the corners up to the top."),
             ("Step 4", "Throw the fly over, match its colored tabs and buckle the corners."),
             ("Step 5", "Stake and tension the guy lines, so the fly doesn't touch the inner tent.")]
    five = "".join(f"<li><b>{step}</b>{esc(what)}</li>" for step, what in steps)
    body = f"""{brand("Product sheet · Tents")}
<div class="top"><div><p class="kicker">New · available November 3, 2026</p><h1>Trailhead 3</h1>
<p class="lede">A three-person dome tent that weighs 1.9 kg and goes up in about three minutes.</p></div>
<div class="art">{dome("#3d7358", "#24493a", "#1d3a2e", "#16261f", "spec")}</div></div>
<table class="spec">{table}</table>
<div class="two"><div><h2 class="first">Two colors</h2><div class="swatches">
<div><i style="background:linear-gradient(135deg,#3d7358,#24493a)"></i><b>Forest green</b></div>
<div><i style="background:linear-gradient(135deg,#e2cfad,#bfa57c)"></i><b>Sand</b></div></div></div>
<div><h2 class="first">In the bag</h2><p>Inner tent and fly · 2 color-coded poles · 10 stakes · 4 guy lines · pole
repair sleeve · stuff sack</p>
<p class="mend"><b>If a pole breaks on a trip:</b> slide the repair sleeve over the break, tape it, and contact us at
contoso.com/help for a new pole. Poles are covered for 5 years.</p></div></div>
<h2>Pitch it in five steps</h2><ol class="five">{five}</ol>{FICTION}"""
    sheets([body], OUTDOOR_CSS, "Trailhead 3 – spec sheet", "Contoso Outdoor · Trailhead 3 · Product sheet",
           out("Documents/Trailhead 3 – spec sheet.pdf"))


# Contoso's own documents: a policy, an event's agenda and two cheat sheets.
POLICY_CSS = """body{font-family:"Avenir Next",system-ui;color:#0f2a4a;font-size:10pt;line-height:1.45}
.head{display:flex;justify-content:space-between;align-items:baseline;padding-bottom:5mm;margin-bottom:10mm;
border-bottom:.8px solid #c9d6e6;font-size:8.5pt;color:#51657d}
.head b{font-size:9.5pt;letter-spacing:.2em;color:#0f2a4a}
.kicker{margin:0 0 2.5mm;font-size:8.5pt;font-weight:700;letter-spacing:.18em;text-transform:uppercase;color:#0f6cbd}
h1{margin:0 0 4mm;font-size:28pt;line-height:1.08;letter-spacing:-.01em}
.lede{margin:0 0 7mm;font-size:12pt;line-height:1.45;color:#51657d;max-width:150mm}
h2{margin:8mm 0 3mm;font-size:13pt;color:#0f6cbd}
h2.first{margin-top:0}
p{margin:0 0 2.6mm}
ul{margin:0 0 2.6mm;padding-left:5mm}li{margin-bottom:1.6mm}
table{width:100%;border-collapse:collapse;margin:1mm 0 3mm}
th{text-align:left;font-size:7.5pt;font-weight:700;letter-spacing:.1em;text-transform:uppercase;color:#51657d;
padding:2mm 3mm;border-bottom:1.5px solid #0f2a4a}
td{padding:2.8mm 3mm;border-bottom:.6px solid #dbe4ef;vertical-align:top}
td:first-child{width:40mm;font-weight:700}
.meta{display:grid;grid-template-columns:repeat(3,1fr);gap:4mm;margin:0 0 8mm}
.meta div{padding:3.5mm 4mm;background:#eef4fb;border-radius:2mm;font-size:8.5pt;color:#51657d}
.meta b{display:block;margin-top:.5mm;font-size:10.5pt;color:#0f2a4a}
.box{padding:5mm 6mm;margin:2mm 0 4mm;border:.8px solid #c9d6e6;border-left:1.3mm solid #0f6cbd;border-radius:1mm}
.box p:last-child{margin:0}
.big{display:flex;gap:8mm;align-items:center;padding:6mm 7mm;background:#0f2a4a;color:#fff;border-radius:3mm;
margin:2mm 0 4mm}
.big strong{font-size:38pt;line-height:1;white-space:nowrap}
.big p{margin:0;font-size:11pt;color:#dbe7f5}
.fiction{margin-top:7mm;font-size:8pt;color:#51657d}
.foot{color:#51657d}
"""


def travel_policy() -> None:
    head = '<div class="head"><b>CONTOSO · FINANZAS</b><span>Política FIN-04 · versión 3</span></div>'
    limits = [("Hotel", "Hasta <b>140 € por noche</b>, con desayuno. Es un límite fijo, sin excepciones: si el hotel "
                        "cuesta más, la diferencia la pagas tú."),
              ("Comidas", "Hasta <b>45 € al día</b> cuando estás de viaje, sin contar el desayuno del hotel."),
              ("Comidas con clientes", "Sin límite fijo, pero en el gasto apunta quién asistió, de qué empresa, el "
                                       "motivo y el código del proyecto (por ejemplo, «TH3 launch»)."),
              ("Coche propio", "<b>0,26 € por kilómetro</b>, más peajes y aparcamiento."),
              ("Avión", "Clase turista. Si el mismo trayecto en tren dura <b>menos de 4 horas</b>, ve en tren."),
              ("Taxi", "Para ir o volver del aeropuerto o de la estación, o <b>después de las 22:00</b>. El resto, "
                       "en transporte público.")]
    rows = "".join(f"<tr><td>{name}</td><td>{what}</td></tr>" for name, what in limits)
    first = f"""{head}<p class="kicker">Política interna</p><h1>Política de viajes y gastos</h1>
<p class="lede">Cómo reservar un viaje, cuánto puedes gastar y cómo recuperar lo que has pagado de tu bolsillo. Se
aplica a todas las personas de Contoso.</p>
<div class="meta"><div>Se aplica desde<b>1 de octubre de 2026</b></div><div>Responsable<b>Equipo de Finanzas</b></div>
<div>Sustituye a<b>Versión 2, de marzo de 2025</b></div></div>
<h2 class="first">Límites</h2><table><tr><th>Gasto</th><th>Qué cubrimos</th></tr>{rows}</table>
<h2>Cómo reservar</h2>
<p>Reserva vuelos, trenes y hoteles en <b>viajes.contoso.com</b>: la agencia ya aplica estos límites y factura
directamente a Contoso, así que no tienes que adelantar nada. Si reservas por tu cuenta, el gasto tiene que cumplir
esta política igualmente.</p>
<div class="box"><p><b>Viajes en grupo.</b> Para eventos y offsites, una sola persona reserva para todo el equipo: así
el hotel y las salas salen más baratos. El límite de 140 € por noche se aplica también a los grupos.</p></div>"""
    second = f"""{head}<h2 class="first">Cómo pedir un reembolso</h2>
<ul><li>Registra el gasto en la aplicación de gastos <b>en un plazo de 30 días</b>, con una foto del recibo o de la
factura.</li>
<li>Indica el proyecto y, si fue una comida con clientes, quién asistió y por qué.</li>
<li>Tu responsable lo aprueba y lo cobras con la siguiente nómina.</li></ul>
<h2>Formación</h2>
<div class="big"><strong>1.500 €</strong><p>A partir de <b>enero de 2027</b>, cada persona tiene un presupuesto de
formación de 1.500 € al año para cursos, certificaciones y conferencias. <b>Hasta 500 € no necesitas aprobación</b>:
basta con registrar el gasto.</p></div>
<p>Por encima de 500 €, pide la aprobación de tu responsable antes de reservar. El presupuesto que no gastes no pasa
al año siguiente.</p>
<h2>Qué no reembolsamos</h2>
<ul><li>El minibar, las multas de tráfico y de aparcamiento.</li>
<li>Las mejoras de clase o de habitación, y los gastos de acompañantes.</li>
<li>Las propinas por encima del 10 %.</li></ul>
<h2>Dudas</h2>
<p>Escribe a <b>viajes@contoso.com</b>. Si algo de esta política no encaja con tu viaje, pregunta antes de
reservar.</p>
<p class="fiction">Contoso es una empresa ficticia y este documento es contenido de ejemplo: sus importes y normas son
inventados.</p>"""
    sheets([first, second], POLICY_CSS, "Política de viajes y gastos",
           "Contoso · Política de viajes y gastos (FIN-04) · versión 3",
           out("Documents/Política de viajes y gastos.pdf"))


SUMMIT_CSS = """body{color:#f0fdff;font-family:"Avenir Next",system-ui;font-size:10pt}
.sheet{overflow:hidden}
.dots{position:absolute;inset:0;
background:radial-gradient(circle,rgba(34,211,238,.5) 1.5px,transparent 1.8px) 0 0/15px 15px;
-webkit-mask-image:radial-gradient(520px 520px at 100% 0%,#000 0%,transparent 70%)}
.beam{position:absolute;inset:0;background:radial-gradient(520px 420px at 0% 100%,rgba(168,85,247,.3),transparent 70%)}
.inner{position:relative}
.top{display:flex;justify-content:space-between;font-size:8pt;font-weight:600;letter-spacing:.16em;
text-transform:uppercase;color:#7dd3fc}
h1{margin:9mm 0 0;font:800 44pt/.95 "Avenir Next Condensed","Avenir Next";letter-spacing:-.01em}
h1 span{color:transparent;-webkit-text-stroke:1.1px #22d3ee}
.when{margin:3.5mm 0 7mm;font-size:11.5pt;color:#bae6fd}
.when b{color:#22d3ee}
table{width:100%;border-collapse:collapse}
td{padding:2mm 0;border-top:.6px solid rgba(186,230,253,.22);vertical-align:top}
td.t{width:18mm;font-weight:700;color:#22d3ee}
td b{display:block;font-size:10.5pt;color:#f0fdff}
td span{display:block;margin-top:.5mm;color:#bae6fd;font-size:8.8pt}
tr.quiet td{color:#8f9abb}
tr.quiet b{display:inline;color:#a8b3d1;font-weight:600}
.rooms{display:grid;grid-template-columns:repeat(3,1fr);gap:2.5mm;margin-top:1.2mm}
.rooms div{padding:2.2mm 3mm;border:.8px solid rgba(34,211,238,.45);border-radius:1.8mm;font-size:9pt;color:#f0fdff}
.rooms i{display:block;margin-bottom:.6mm;font-style:normal;font-size:7pt;font-weight:700;letter-spacing:.14em;
text-transform:uppercase;color:#c084fc}
.keynote b{color:#f0fdff}.keynote em{font-style:normal;color:#c084fc;font-weight:700}
.notes{display:flex;justify-content:space-between;margin-top:5mm;padding-top:3mm;
border-top:.6px solid rgba(186,230,253,.22);font-size:9pt;font-weight:600;color:#7dd3fc}
.fiction{margin-top:2.5mm;font-size:7.5pt;color:#8f9abb}
.foot{color:#8f9abb}
"""


def summit_agenda() -> None:
    def rooms(*talks: str) -> str:
        return '<div class="rooms">' + "".join(f"<div><i>Room {room}</i>{esc(talk)}</div>"
                                               for room, talk in zip("ABC", talks, strict=True)) + "</div>"

    agenda = [("8:00", "Doors open and coffee", "", True),
              ("9:00", "Welcome", "Data Summit Madrid 2026: the day in five minutes.", False),
              ("9:15", "<em>Opening keynote</em> · How Contoso Outdoor handles returns with an AI agent",
               "Lucía Romero, Head of Customer Experience, Contoso Outdoor.", False),
              ("9:30", "One copy of your data: Microsoft Fabric from end to end",
               "Ingestion, OneLake, Power BI in Direct Lake mode and Copilot in Fabric, on one platform.", False),
              ("10:30", "Coffee break", "", True),
              ("11:00", "Breakouts", rooms("OneLake shortcuts and mirroring", "Building agents with Microsoft Foundry",
                                           "GitHub Copilot agent mode, hands-on"), False),
              ("12:00", "Breakouts", rooms("Real-Time Intelligence and KQL", "RAG and evaluations: measuring an agent",
                                           "MCP servers: giving agents tools"), False),
              ("13:00", "Lightning talks", "Five ten-minute talks from teams that shipped something with data and AI "
                                           "this year.", False),
              ("14:00", "Standing lunch", "", True),
              ("15:00", "Live demos", "Microsoft Fabric, Microsoft Foundry and GitHub Copilot, built live on stage.",
               False),
              ("16:00", "Coffee break", "", True),
              ("16:30", "Panel: AI in production, what we learned", "Four teams on evaluations, costs, rate limits "
                                                                    "and what they'd do differently.", False),
              ("17:30", "Closing remarks", "", True)]
    rows = []
    for when, title, detail, quiet in agenda:
        title = title if "<em>" in title else esc(title)
        if detail and not detail.startswith("<div"):
            detail = f"<span>{esc(detail)}</span>"
        cls = ' class="quiet"' if quiet else (' class="keynote"' if "<em>" in title else "")
        rows.append(f'<tr{cls}><td class="t">{when}</td><td><b>{title}</b>{detail}</td></tr>')
    body = f"""<div class="dots"></div><div class="beam"></div><div class="inner">
<div class="top"><span>Contoso Outdoor presents</span><span>Agenda · version 2</span></div>
<h1>DATA SUMMIT <span>MADRID</span></h1>
<p class="when"><b>Thursday, November 12, 2026</b> · Doors open 8:00 · First session 9:00</p>
<table>{"".join(rows)}</table>
<div class="notes"><span>900+ registered · Free</span><span>Coffee and lunch by Fourth Coffee</span></div>
<p class="fiction">Vegetarian and gluten-free menus: tell us when you register. Contoso Outdoor and Fourth Coffee are
fictitious companies, and this agenda is sample content.</p></div>"""
    sheets([body], SUMMIT_CSS, "Data Summit Madrid 2026 – agenda", "Data Summit Madrid 2026 · Thursday, November 12",
           out("Documents/Data Summit Madrid 2026 – agenda.pdf"), "#050510")


KQL_KEYWORDS = {"where", "project", "summarize", "by", "extend", "join", "kind", "inner", "on", "take", "top", "render",
                "make-series", "default", "from", "to", "step", "between", "and", "or", "has", "contains", "with",
                "asc", "desc", "timechart"}
KQL_FUNCTIONS = {"ago", "bin", "count", "countif", "dcount", "percentile", "avg", "max", "min", "datetime", "now",
                 "series_decompose_anomalies"}
KQL_TABLES = {"AppRequests", "AppDependencies", "ToolCalls"}


def kql(source: str) -> str:
    """KQL as HTML, colored as the Azure Data Explorer web UI roughly does."""
    parts, last = [], 0
    for match in re.finditer(r'//[^\n]*|"[^"]*"|\d[\w.]*|[A-Za-z_][\w-]*|\|', source):
        token = match.group()
        kind = ("c" if token.startswith("//") else "s" if token.startswith('"') else "n" if token[0].isdigit()
                else "p" if token == "|" else "k" if token in KQL_KEYWORDS else "f" if token in KQL_FUNCTIONS
                else "t" if token in KQL_TABLES else "")
        parts.append(esc(source[last:match.start()]))
        parts.append(f'<span class="{kind}">{esc(token)}</span>' if kind else esc(token))
        last = match.end()
    parts.append(esc(source[last:]))
    return "".join(parts)


KQL_CARDS = [
    ("1 · The shape of a query", "A table, then operators one after another, each after a pipe.",
     """AppRequests
| where TimeGenerated > ago(1d)
| where Name == "POST /api/messages"
| project TimeGenerated, OperationId, DurationMs, Success
| take 10"""),
    ("2 · Count by time", "The 429s from the model on October 13, in 5-minute bins.",
     """AppDependencies
| where TimeGenerated between (datetime(2026-10-13 10:30) .. datetime(2026-10-13 12:00))
| where ResultCode == "429"
| summarize Throttled = count() by bin(TimeGenerated, 5m)
| render timechart"""),
    ("3 · Percentiles", "Latency of the chat endpoint. CX-147 alerts when p95 goes over 8 s.",
     """AppRequests
| where TimeGenerated > ago(7d)
| where Name == "POST /api/messages"
| summarize p50 = percentile(DurationMs, 50),
            p95 = percentile(DurationMs, 95) by bin(TimeGenerated, 1h)
| where p95 > 8000"""),
    ("4 · Errors in a window", "The other CX-147 alert: more than 10 failed requests in 5 minutes.",
     """AppRequests
| where TimeGenerated > ago(1h)
| summarize Errors = countif(Success == false) by bin(TimeGenerated, 5m)
| where Errors > 10"""),
    ("5 · Join two tables", "Which tools the assistant calls, and how long those requests take.",
     """ToolCalls  // e.g. lookup_order, create_return, book_pickup
| where TimeGenerated > ago(7d)
| join kind=inner (AppRequests | project OperationId, DurationMs) on OperationId
| summarize Calls = count(), p95 = percentile(DurationMs, 95) by Tool
| top 5 by Calls desc"""),
    ("6 · Spot anomalies", "Requests per hour as a series, and the hours that don't fit the pattern.",
     """AppRequests
| make-series Requests = count() default = 0
    on TimeGenerated from ago(14d) to now() step 1h
| extend Anomalies = series_decompose_anomalies(Requests)
| render timechart"""),
]

KQL_OPERATORS = [("where", "Keeps the rows that match a condition."),
                 ("project", "Chooses, renames or computes columns."),
                 ("extend", "Adds computed columns and keeps the rest."),
                 ("summarize", "Aggregates: count(), avg(), percentile()… by groups."),
                 ("bin()", "Rounds values down into buckets, as 5m or 1h."),
                 ("join", "Combines rows of two tables on a common column."),
                 ("take", "Any n rows: a quick look at the data."),
                 ("top", "The first n rows, sorted by a column."),
                 ("make-series", "A series per group, with empty bins filled in."),
                 ("render", "Draws the result: timechart, barchart, piechart…")]

KQL_CSS = """body{font-family:"Avenir Next",system-ui;color:#1b1f2a;font-size:9.5pt;line-height:1.4}
.head{display:flex;justify-content:space-between;align-items:flex-end;padding-bottom:4mm;margin-bottom:5mm;
border-bottom:1.5px solid #1b1f2a}
.head h1{margin:0;font-size:24pt;line-height:1;letter-spacing:-.01em}
.head h1 i{font-style:normal;color:#0b57d0}
.head span{font-size:8.5pt;color:#5d6475;text-align:right}
.lede{margin:0 0 5mm;font-size:10pt;color:#5d6475}
.cards{display:grid;grid-template-columns:1fr 1fr;gap:3.5mm}
.card{border:.8px solid #dfe3ea;border-radius:2.5mm;padding:3.2mm 3.6mm}
.card h2{margin:0 0 .8mm;font-size:10pt}
.card p{margin:0 0 2.2mm;font-size:8.3pt;color:#5d6475}
pre{margin:0;padding:2.6mm 3mm;background:#f5f7fb;border-radius:1.5mm;font:7.1pt/1.5 "SF Mono",Menlo,monospace;
white-space:pre-wrap;color:#1b1f2a}
.k{color:#0b57d0;font-weight:600}.f{color:#8430ce}.s{color:#b3261e}.n{color:#0f7b6c}.t{font-weight:700}
.p{color:#9aa1ae}.c{color:#7c8494;font-style:italic}
.wide{grid-column:1/-1}
.ops{columns:2;column-gap:7mm}
.ops div{display:flex;gap:3mm;padding:1.3mm 0;border-bottom:.6px solid #e8ebf0;break-inside:avoid;font-size:8.3pt}
.ops b{flex:none;width:23mm;font:600 7.8pt "SF Mono",Menlo,monospace;color:#0b57d0}
.where{display:flex;flex-wrap:wrap;gap:1.8mm;margin-top:1mm}
.where span{padding:1mm 2.6mm;border-radius:99px;background:#eef3fd;color:#0b3d91;font-size:8pt;font-weight:600}
.small{margin:3.5mm 0 0;font-size:7.6pt;color:#7c8494}
.list div{padding:1.7mm 0;border-bottom:.6px solid #e8ebf0;font-size:8.3pt;color:#5d6475}
.list div:last-child{border:0;padding-bottom:0}
.list b{display:block;margin-bottom:.3mm;font-size:9pt;color:#1b1f2a}
.list code{font:600 7.8pt "SF Mono",Menlo,monospace;color:#0b57d0}
.list b code{font-size:8.3pt}
.foot{color:#7c8494}
"""


KQL_HABITS = [("Filter on time first", "Put <code>where TimeGenerated &gt; ago(…)</code> right after the table: "
                                       "every step after it reads less."),
              ("<code>has</code>, not <code>contains</code>", "<code>has</code> looks for a whole term and uses the "
                                                              "index; <code>contains</code> reads every string."),
              ("Project early", "Keep only the columns you need, as soon as you can."),
              ("<code>==</code> before <code>=~</code>", "Case-sensitive comparisons are faster: use "
                                                         "<code>=~</code> only when the case really varies."),
              ("<code>take</code> before you query", "<code>take 10</code> shows a table's shape in a second, before "
                                                    "you query all of it.")]

KQL_TIME = [("ago(7d)", "Seven days before now."),
            ("now()", "This moment. KQL's times are in UTC."),
            ("startofday(now())", "Midnight today, UTC."),
            ("bin(TimeGenerated, 1h)", "Hourly buckets, for summarize."),
            ("datetime(2026-10-13 10:30)", "A fixed moment."),
            ("1d · 2h · 30m · 10s · 100ms", "Timespans: days, hours, minutes, seconds, milliseconds.")]


def kql_cheat_sheet() -> None:
    cards = "".join(f'<div class="card"><h2>{esc(title)}</h2><p>{esc(why)}</p><pre>{kql(query)}</pre></div>'
                    for title, why, query in KQL_CARDS)
    operators = "".join(f"<div><b>{esc(name)}</b><span>{esc(what)}</span></div>" for name, what in KQL_OPERATORS)
    habits = "".join(f"<div><b>{name}</b>{why}</div>" for name, why in KQL_HABITS)
    times = "".join(f"<div><b><code>{esc(name)}</code></b>{esc(what)}</div>" for name, what in KQL_TIME)
    places = ["Azure Data Explorer", "Azure Monitor Logs", "Application Insights", "Microsoft Sentinel",
              "Microsoft Defender XDR · advanced hunting", "Microsoft Fabric · Real-Time Intelligence (Eventhouse)"]
    where = "".join(f"<span>{esc(place)}</span>" for place in places)
    head = """<div class="head"><h1>KQL <i>cheat sheet</i></h1><span>Contoso · Customer Experience data<br>
For the returns assistant's logs · October 2026</span></div>"""
    first = f"""{head}
<p class="lede">Kusto Query Language reads like a pipeline: start from a table and pass its rows through one operator
after another. Six queries we use every week, from the simplest to the most useful.</p>
<div class="cards">{cards}</div>"""
    second = f"""{head}
<div class="cards"><div class="card wide"><h2>Ten operators to start with</h2><div class="ops">{operators}</div></div>
<div class="card"><h2>Five habits for fast queries</h2><div class="list">{habits}</div></div>
<div class="card"><h2>Time, the KQL way</h2><div class="list">{times}</div></div>
<div class="card wide"><h2>Where KQL runs</h2><p>The same language, in every one of these:</p>
<div class="where">{where}</div></div></div>
<p class="small">AppRequests and AppDependencies are Application Insights tables in a Log Analytics workspace.
ToolCalls is the assistant's own table, made up for this sheet: in a Log Analytics workspace a custom table's name ends
in _CL. The habits come from Microsoft Learn's best practices for KQL queries. Contoso is a fictitious company.</p>"""
    sheets([first, second], KQL_CSS, "KQL cheat sheet", "Contoso · KQL cheat sheet · Javier Molina",
           out("Documents/KQL cheat sheet.pdf"))


PROMPTS_CSS = """body{font-family:"Avenir Next",system-ui;color:#1e1b2e;font-size:9.8pt;line-height:1.42}
.head{display:flex;justify-content:space-between;align-items:baseline;font-size:8.5pt;color:#6b6585;
padding-bottom:4mm;margin-bottom:7mm;border-bottom:.8px solid #e3defa}
.head b{letter-spacing:.2em;color:#6d28d9}
h1{margin:0 0 3mm;font-size:27pt;line-height:1.05;letter-spacing:-.01em}
.lede{margin:0 0 6mm;font-size:11.5pt;color:#6b6585;max-width:155mm}
h2{margin:6mm 0 2.5mm;font-size:12.5pt;color:#6d28d9}
h2.first{margin-top:0}
p{margin:0 0 2.4mm}
.four{display:grid;grid-template-columns:repeat(4,1fr);gap:3mm;margin-bottom:3mm}
.four div{padding:3.2mm 3.5mm;border-radius:2.2mm;font-size:8.8pt;color:#3b3552}
.four b{display:block;margin-bottom:.8mm;font-size:10pt;color:#1e1b2e}
.g{background:#ede9fe}.c{background:#dbeafe}.e{background:#dcfce7}.s{background:#ffedd5}
.example{padding:4mm 5mm;border:.8px solid #e3defa;border-radius:2.5mm;font-size:10.5pt;line-height:1.6}
.example mark{padding:.4mm 1mm;border-radius:1mm;color:inherit}
mark.g{background:#ede9fe}mark.c{background:#dbeafe}mark.e{background:#dcfce7}mark.s{background:#ffedd5}
.source{margin-top:2mm;font-size:8pt;color:#6b6585}
.six{display:grid;grid-template-columns:1fr 1fr;gap:3mm}
.six div{padding:3mm 3.6mm;border:.8px solid #e3defa;border-radius:2.2mm}
.six i{display:block;margin-bottom:1mm;font-style:normal;font-size:7.5pt;font-weight:700;letter-spacing:.14em;
text-transform:uppercase;color:#6d28d9}
.six q{font-size:9.4pt}
.tips{display:grid;grid-template-columns:1fr 1fr;gap:6mm}
ul{margin:0;padding-left:4.5mm}li{margin-bottom:1.5mm}
code{font:8.6pt "SF Mono",Menlo,monospace;background:#f3f0ff;padding:.3mm 1mm;border-radius:1mm}
.fiction{margin-top:5mm;font-size:7.8pt;color:#6b6585}
h1.sub{margin:0 0 6mm;font-size:21pt}
.keys div{display:flex;gap:3mm;padding:1.6mm 0;border-bottom:.6px solid #e3defa}
.keys div:first-child{padding-top:0}
.keys code{flex:none;width:23mm;padding:0;background:none;color:#6d28d9;font-weight:600}
.foot{color:#6b6585}
"""


def prompt_cheat_sheet() -> None:
    examples = [("Teams", "Resume la reunión: decisiones, tareas con su responsable y fecha, y lo que quedó "
                          "pendiente."),
                ("Outlook", "Resume los correos sin leer de esta semana sobre la Trailhead 3 y dime cuáles necesitan "
                            "una respuesta mía."),
                ("Word", "Convierte estas notas en un informe de una página, con un resumen arriba y los próximos "
                         "pasos al final."),
                ("Excel", "¿Qué motivos de devolución han crecido más este trimestre? Muéstralo en un gráfico de "
                          "barras."),
                ("PowerPoint", "Crea una presentación de seis diapositivas a partir de este documento, con notas del "
                               "orador en cada una."),
                ("GitHub Copilot", "/tests Escribe tests con pytest para la función seleccionada: el caso normal, "
                                   "los límites y los errores.")]
    six = "".join(f"<div><i>{app}</i><q>{esc(prompt)}</q></div>" for app, prompt in examples)
    shortcuts = [("/explain", "Explica el código seleccionado."),
                 ("/fix", "Propone un arreglo para lo seleccionado."),
                 ("/tests", "Escribe tests para lo seleccionado."),
                 ("/new", "Crea un proyecto nuevo a partir de una descripción."),
                 ("#file", "Añade un archivo concreto como contexto."),
                 ("@workspace", "Responde sobre todo el proyecto, no solo el archivo abierto.")]
    keys = "".join(f"<div><code>{key}</code><span>{esc(what)}</span></div>" for key, what in shortcuts)
    head = '<div class="head"><b>CONTOSO · PLATAFORMA</b><span>Para todo Contoso · septiembre de 2026</span></div>'
    first = f"""{head}
<h1>Chuleta de prompts para Copilot</h1>
<p class="lede">Cómo pedirle las cosas a Microsoft 365 Copilot y a GitHub Copilot para que acierte a la primera, y qué
hacer cuando no acierta.</p>
<h2 class="first">Los cuatro elementos de un buen prompt</h2>
<div class="four"><div class="g"><b>Objetivo</b>Qué quieres que haga.</div>
<div class="c"><b>Contexto</b>Por qué lo necesitas y para quién es.</div>
<div class="e"><b>Expectativas</b>Formato, tono, idioma y extensión.</div>
<div class="s"><b>Fuente</b>Qué archivos, correos o chats debe usar.</div></div>
<div class="example"><mark class="g">Redacta un correo para Tom Becker, de Fabrikam Logistics,</mark>
<mark class="c">para confirmar que el pedido PO 5518, las varillas de la Trailhead 3, llega a nuestro almacén de
Zaragoza el 13 de octubre: sin ellas no podemos empezar el montaje.</mark> <mark class="e">En inglés, cordial y en menos
de 120 palabras, con una lista de lo que necesitamos que nos confirme.</mark> <mark class="s">Usa el hilo «PO 5518» de
Outlook y la hoja de pedidos de Excel.</mark></div>
<p class="source">Los cuatro elementos son los de la guía pública de Microsoft para escribir prompts para Copilot.</p>
<h2>Seis prompts para empezar</h2><div class="six">{six}</div>"""
    second = f"""{head}
<h1 class="sub">Después del primer prompt</h1>
<div class="tips"><div><h2 class="first">Itera, no empieces de cero</h2>
<p>Si la primera respuesta no es lo que buscabas, pide cambios en la misma conversación: «más corto», «en una tabla»,
«añade las fechas», «en tono más formal».</p></div>
<div><h2 class="first">Revisa siempre</h2><p>Copilot puede equivocarse. Comprueba cifras, nombres y fechas antes de
enviar nada, y no pegues datos de clientes fuera de las herramientas aprobadas.</p></div></div>
<h2>Con GitHub Copilot</h2>
<div class="tips"><ul>
<li>Pon las normas del repositorio en <code>.github/copilot-instructions.md</code>: Copilot las lee en cada
petición.</li>
<li>Abre los archivos relevantes: Copilot los usa como contexto.</li>
<li>Pide los tests a la vez que el código.</li>
<li>En modo agente, revisa el diff antes de aceptar los cambios.</li></ul>
<div><div class="keys">{keys}</div>
<p class="source">Comandos de Copilot Chat en Visual Studio Code. La lista completa está en docs.github.com.</p></div>
</div>
<h2>Con Microsoft 365 Copilot</h2>
<div class="tips"><ul><li>Di dónde buscar: «en el hilo PO 5518», «en la reunión del martes».</li>
<li>Pide que cite: cada dato, con el correo o el documento del que sale.</li></ul>
<ul><li>Guarda en Copilot Prompt Gallery los prompts que te funcionen y compártelos con tu equipo.</li>
<li>Si la respuesta es larga, pide primero un resumen de tres líneas.</li></ul></div>
<p class="fiction">Contoso y Fabrikam son empresas ficticias: los nombres y pedidos de los ejemplos son
inventados.</p>"""
    sheets([first, second], PROMPTS_CSS, "Chuleta de prompts para Copilot", "Contoso · Chuleta de prompts para Copilot",
           out("Documents/Chuleta de prompts para Copilot.pdf"))


# MARK: Word and RTF


def word(body: str) -> str:
    """HTML that textutil turns into a Word document, in make_demo's documents()' style."""
    return f"""<!doctype html><html><head><meta charset="utf-8"><style>body{{font-family:'Helvetica Neue';
font-size:11pt}} h1{{font-size:20pt}} h2{{font-size:13pt;color:#0f4c81}} th{{background:#e8eef6;text-align:left}}
</style></head><body>{body}</body></html>"""


def table(header: list[str], rows: list[tuple]) -> str:
    head = "".join(f"<th>{cell}</th>" for cell in header)
    body = "".join("<tr>" + "".join(f"<td>{cell}</td>" for cell in row) + "</tr>" for row in rows)
    return f'<table border="1" cellpadding="6" cellspacing="0"><tr>{head}</tr>{body}</table>'


PRESS_RELEASE = """<p><b>NOTA DE PRENSA · BORRADOR v2 · NO DIFUNDIR</b><br><i>Embargo: martes 3 de noviembre de 2026,
9:00 CET</i></p>
<h1>Contoso Outdoor presenta Trailhead 3: tres plazas, 1,9 kg y montada en tres minutos</h1>
<p><b>La nueva tienda de campaña de la marca pesa 300 gramos menos que su predecesora, tiene dos puertas y dos ábsides
y llega a las tiendas y a contoso.com el 3 de noviembre, por 249 euros.</b></p>
<p>Madrid, 3 de noviembre de 2026. Contoso Outdoor, fabricante español de material de montaña y acampada, presenta hoy
Trailhead 3, una tienda de campaña tipo iglú para tres personas que pesa 1,9 kg y se monta en unos tres minutos
gracias a sus varillas de aluminio con código de colores.</p>
<p>Trailhead 3 nace de los datos. Al analizar las devoluciones con Microsoft Fabric, el equipo de Contoso Outdoor vio
que el 31 % de las devoluciones de septiembre eran de productos que habían llegado dañados, muchas veces con una
varilla rota. Por eso la nueva tienda lleva una funda de reparación para varillas y un embalaje rediseñado, y sus
varillas tienen cinco años de garantía.</p>
<p>«Cuando un cliente nos escribe porque su tienda ha llegado con una varilla rota, nuestro asistente abre la
devolución y le envía otra en minutos. Con Trailhead 3 queremos que ese mensaje no llegue nunca», explica Lucía
Romero, directora de Experiencia de Cliente de Contoso Outdoor.</p>
<h2>Principales características</h2>
<ul><li>Tres plazas, 1,9 kg y 46 × 15 cm plegada.</li>
<li>Dos puertas y dos ábsides, para que nadie tenga que pasar por encima de nadie.</li>
<li>Montaje en unos tres minutos, con varillas de aluminio con código de colores.</li>
<li>Doble techo de poliéster reciclado.</li>
<li>Funda de reparación para varillas incluida. Dos años de garantía, cinco en las varillas.</li></ul>
<h2>Precio y disponibilidad</h2>
<p>Trailhead 3 estará disponible a partir del 3 de noviembre de 2026 en las tiendas de Contoso Outdoor y en
contoso.com, en verde bosque y en arena, por 249 euros.</p>
<h2>Sobre Contoso Outdoor</h2>
<p>Contoso Outdoor diseña y vende material de montaña y acampada. Tiene su sede en Madrid y envía todos sus pedidos
desde su almacén de Zaragoza.</p>
<p><b>Contacto de prensa:</b> Marta Gil · prensa@contoso.com</p>
<h2>Notas para Lucía (borrar antes de enviar)</h2>
<ul><li>Legal pide quitar «la más ligera del mercado»: no lo podemos demostrar. Lo he cambiado por «300 gramos menos
que su predecesora».</li>
<li>¿Mantenemos la frase del asistente? Si preferís guardarla para la keynote del 12 de noviembre, la quito.</li>
<li>Las fotos de producto llegan el 28 de octubre.</li>
<li>Necesito tu revisión antes del viernes, por favor.</li></ul>
<p><i>Marta Gil, 14 de octubre. Contoso Outdoor es una empresa ficticia y esta nota de prensa es contenido de
ejemplo.</i></p>"""

OFFSITE_AGENDA = f"""<h1>Offsite de Experiencia de Cliente · Lisboa</h1>
<p><b>Del miércoles 21 al viernes 23 de octubre de 2026</b> · 14 personas · Organiza: Ana Castillo</p>
<h2>Vuelos</h2>
{table(["Vuelo", "Día", "Salida", "Llegada"],
       [("BY 471", "Miércoles 21 de octubre", "Madrid T4, 08:35", "Lisboa T1, 08:55"),
        ("BY 476", "Viernes 23 de octubre", "Lisboa T1, 19:40", "Madrid T4, 21:55")])}
<p>Reserva MT-4Q7ZK, para todo el equipo. Lisboa va una hora por detrás de Madrid: las horas de Lisboa son
locales.</p>
<h2>Hotel</h2>
<p>Un hotel en la Baixa, con desayuno y una sala para las sesiones, por menos de 140\u00a0€ por noche, como pide la
política de viajes. Habitaciones individuales.</p>
<h2>Miércoles 21</h2>
{table(["Hora", "Qué", "Quién"],
       [("10:30", "Llegada al hotel y café", ""),
        ("11:00–13:00", "Bienvenida y resultados del piloto", "Lucía"),
        ("13:00", "Comida", ""),
        ("14:30–17:30", "Taller: el asistente en 2027", "Ana y Sara"),
        ("Tarde", "Libre", "")])}
<h2>Jueves 22</h2>
{table(["Hora", "Qué", "Quién"],
       [("9:00–13:00", "Estrategia 2027", "Lucía"),
        ("13:00", "Comida", ""),
        ("14:30–17:00", "Visita a la futura oficina de Lisboa", "Lucía y Pablo"),
        ("20:30", "Cena en Alfama, junto al mirador", "Todos")])}
<h2>Viernes 23</h2>
{table(["Hora", "Qué", "Quién"],
       [("9:30–12:30", "Plan del primer trimestre y retro del offsite", "Ana"),
        ("12:30", "Check-out: las maletas se quedan en el hotel", ""),
        ("17:30", "Taxi al aeropuerto, desde el hotel", ""),
        ("19:40", "Vuelo BY 476 a Madrid", "")])}
<h2>A tener en cuenta</h2>
<ul><li>Cena del jueves: somos 14, con 3 menús vegetarianos y 1 sin gluten. Si tienes alguna otra alergia o
intolerancia, dímelo antes del martes 20 a las 12:00.</li>
<li>Del hotel al restaurante hay 10 minutos cuesta arriba: zapato cómodo.</li>
<li>Los gastos siguen la política de viajes (FIN-04): comidas hasta 45\u00a0€ al día. El taxi al aeropuerto entra.</li>
<li>La oficina de Lisboa abre en marzo de 2027. Esperamos ser 40 personas allí a finales de 2027.</li></ul>
<p><i>Contoso es una empresa ficticia y esta agenda es contenido de ejemplo.</i></p>"""

KEYNOTE_SCRIPT_SLIDES = [
    ("1", "How Contoso Outdoor handles returns with an AI agent", "0:30", "0:30",
     "Good morning! I'm Lucía, and I look after customer experience at Contoso Outdoor. We make tents. And sometimes "
     "a tent arrives with a broken pole. This is the story of what happens next."),
    ("2", "Eleven minutes, by hand", "1:30", "2:00",
     "A year ago, a person handled every return: find the order, read the policy, open the return, email a label. "
     "Eleven minutes, every time, for almost the same answer. And September was our worst month."),
    ("3", "Now 62 % resolve on their own", "1:30", "3:30",
     "Today 62 % of those conversations resolve on their own, in about four minutes, and customers rate them 4.5 out "
     "of 5. The rest go to a person, with the whole case already written up."),
    ("4", "Three things made it work", "2:00", "5:30",
     "One: our data in one place, in OneLake, with Microsoft Fabric: orders, returns and the policies. Two: an agent "
     "with tools, built with Microsoft Foundry: it can look up an order, open a return and book a pickup. Three: "
     "evaluations. 120 golden questions we run before every release."),
    ("5", "Live demo: a tent with a broken pole", "3:30", "9:00",
     "Let's do it live. I'm Alex, and my Trailhead 2 arrived with a broken pole. Watch the tool calls: lookup_order "
     "finds order 48213, the policy says damaged items come back free, create_return opens RMA-20931, and "
     "book_pickup offers a courier slot. If the demo fails, play Screen Recording 2026-10-14 at 16.31.52 from the "
     "desktop: it's the same conversation."),
    ("6", "What we learned", "1:30", "10:30",
     "Start with the boring questions: they're most of the volume. Write the evaluations before the prompt. Plan for "
     "rate limits: on October 13 we got 429s for twenty minutes, and retries with backoff fixed it. And hand over "
     "to a person as soon as a customer is upset."),
    ("7", "What the data changed: Trailhead 3", "1:00", "11:30",
     "And the data changed the product. 31 % of September's returns had arrived damaged, many with a broken pole. "
     "So the team redesigned the tent: Trailhead 3. Three people, 1.9 kilos, color-coded aluminum poles and a repair "
     "sleeve in the bag. It's been in stores since November 3."),
    ("8", "Thank you", "0:30", "12:00",
     "Thank you! Questions at the coffee break: I'll be next to the Fourth Coffee stand."),
]

KEYNOTE_SCRIPT = f"""<h1>Opening keynote · Data Summit Madrid</h1>
<p><b>How Contoso Outdoor handles returns with an AI agent</b><br>Lucía Romero · Thursday, November 12, 2026 · 9:15 ·
12 minutes</p>
<p><i>Version 3 · October 15. Since the rehearsal: the architecture slide is gone, the live demo stays, and we end
with the Trailhead 3 story.</i></p>
<h2>Timing</h2>
{table(["#", "Slide", "Time", "Ends at"],
       [(number, title, length, ends) for number, title, length, ends, _ in KEYNOTE_SCRIPT_SLIDES]
       + [("", "<b>Total</b>", "<b>12:00</b>", "")])}
<h2>Script</h2>
{"".join(f"<p><b>{number} · {title}</b> ({length})<br>{said}</p>"
         for number, title, length, _, said in KEYNOTE_SCRIPT_SLIDES)}
<h2>On the day</h2>
<ul><li>Sara resets the demo order before 9:00 and sits in the front row with the backup laptop.</li>
<li>Backup video: Screen Recording 2026-10-14 at 16.31.52, on the desktop of both laptops.</li>
<li>Clicker, water, and the timer on the stage monitor at 12:00.</li></ul>
<p><i>Contoso Outdoor is a fictitious company, and this script is sample content.</i></p>"""

SECURITY_MINUTES = f"""<h1>Acta · Comité de seguridad</h1>
<p><b>Viernes 16 de octubre de 2026, 12:00–13:00</b> · Sala Ártico y Teams</p>
<p><b>Asisten:</b> Marc Puig (preside), Lucía Romero, Sara Navarro y Javier Molina. <b>Excusa su asistencia:</b> Pablo
Serrano.</p>
<h2>1. Simulacro de phishing de septiembre</h2>
<p>El 18 % de la plantilla hizo clic en el enlace (en junio, el 31 %) y el 64 % lo notificó con el botón «Notificar
phishing». El próximo simulacro será en diciembre. Marc recuerda la guía «Phishing – what to do if you clicked».</p>
<h2>2. Clave de staging expuesta</h2>
<p>La clave de API del entorno de staging del asistente apareció en un log. Se rotó el jueves 15 y no hay indicios de
uso. Para que no vuelva a pasar, el asistente se autenticará con identidad administrada (ADR 0007), sin claves en la
configuración: CX-148, Marc, antes del 30 de octubre.</p>
<h2>3. Servidor MCP de pedidos</h2>
<p>El servidor MCP de pedidos solo expone tres herramientas: lookup_order, create_return y book_pickup. Mínimo
privilegio: leer pedidos y crear devoluciones y recogidas, nada más. Sara revisará sus permisos antes de que se
incorpore el equipo de ventas, el 2 de noviembre (CX-150).</p>
<h2>4. Políticas de GitHub Copilot</h2>
<p>Antes de la primera ola del despliegue, el 19 de octubre, quedan configuradas en la organización: bloquear las
sugerencias que coincidan con código público, excluir del contexto de Copilot los archivos de configuración con
secretos y exigir la revisión de una persona en todos los pull requests. Responsable: Marc.</p>
<h2>5. Alertas del asistente</h2>
<p>CX-147, hecho: hay alertas cuando el p95 de latencia supera 8 s o cuando hay más de 10 errores en 5 minutos. Sara
las probó ayer.</p>
<h2>Acciones</h2>
{table(["Acción", "Responsable", "Fecha"],
       [("Identidad administrada en el asistente (CX-148)", "Marc", "30 de octubre"),
        ("Revisión de permisos del servidor MCP (CX-150)", "Sara", "Antes del 2 de noviembre"),
        ("Políticas de GitHub Copilot en la organización", "Marc", "Antes del 19 de octubre"),
        ("Simulacro de phishing", "Marc", "Diciembre")])}
<p>Próxima reunión: viernes 20 de noviembre, a las 12:00, en la Sala Ártico.</p>
<p><i>Contoso es una empresa ficticia y esta acta es contenido de ejemplo.</i></p>"""

AI_GLOSSARY = [
    ("Agente", "Un sistema que combina un modelo con instrucciones y herramientas para completar tareas: no solo "
               "responde, también consulta datos y actúa, como abrir una devolución."),
    ("Ajuste fino (fine-tuning)", "Volver a entrenar un modelo con ejemplos propios para que aprenda un estilo o una "
                                  "tarea concreta. Suele ser el último recurso, después de mejorar el prompt y de "
                                  "usar RAG."),
    ("Alucinación", "Una respuesta que suena convincente pero es falsa o no se apoya en las fuentes. La fundamentación "
                    "y las evaluaciones sirven para reducirlas."),
    ("Copilot", "La marca de los asistentes de IA de Microsoft: Microsoft 365 Copilot en las aplicaciones de Office, "
                "GitHub Copilot para programar y Copilot en Microsoft Fabric para trabajar con datos, entre otros."),
    ("Embedding", "Un vector de números que representa el significado de un texto, una imagen o un audio. Contenidos "
                  "parecidos tienen vectores cercanos: es la base de la búsqueda semántica."),
    ("Evaluaciones", "Medir la calidad de un modelo o de un agente con un conjunto de preguntas de referencia (un "
                     "«golden set») antes de publicarlo. Microsoft Foundry incluye evaluadores de fundamentación, "
                     "relevancia, coherencia y seguridad."),
    ("Fundamentación (grounding)", "Que la respuesta se base en fuentes concretas, como la política de devoluciones, "
                                   "y no solo en lo que el modelo aprendió al entrenarse."),
    ("Guardrails", "Filtros y reglas que limitan lo que un modelo o un agente puede decir o hacer: filtros de "
                   "contenido, temas prohibidos, herramientas permitidas o aprobación de una persona."),
    ("Límite de tasa (429)", "El máximo de peticiones o de tokens por minuto que admite un modelo desplegado. Al "
                             "superarlo, la API responde «429 Too Many Requests»: hay que reintentar con espera "
                             "exponencial."),
    ("LLM", "Modelo de lenguaje grande (large language model): un modelo entrenado con enormes cantidades de texto, "
            "capaz de entender y generar lenguaje."),
    ("MCP (Model Context Protocol)", "Un protocolo abierto, presentado por Anthropic en noviembre de 2024, para "
                                     "conectar modelos y agentes con herramientas y datos de forma estándar."),
    ("Microsoft Foundry", "La plataforma de Azure para crear aplicaciones y agentes de IA: modelos, agentes, "
                          "herramientas, evaluaciones y supervisión en un mismo sitio. Se llamaba Azure AI Foundry "
                          "hasta Ignite 2025."),
    ("Multimodal", "Un modelo que trabaja con varios tipos de contenido a la vez: texto, imágenes, audio o vídeo."),
    ("Prompt", "Las instrucciones y el contexto que se le dan a un modelo. Un buen prompt tiene objetivo, contexto, "
               "expectativas y fuente."),
    ("RAG (retrieval-augmented generation)", "Generación aumentada por recuperación: antes de responder, el sistema "
                                             "busca los fragmentos relevantes en documentos propios y se los pasa al "
                                             "modelo como contexto."),
    ("SLM", "Modelo de lenguaje pequeño (small language model), como los de la familia Phi de Microsoft: más barato y "
            "rápido, y capaz de funcionar en un portátil o en un móvil."),
    ("Temperatura", "Un parámetro que controla cuánto varían las respuestas. Con valores bajos son más predecibles; "
                    "con valores altos, más creativas."),
    ("Token", "La unidad en la que un modelo lee y escribe texto. En inglés, un token equivale de media a unos cuatro "
              "caracteres. Los límites y los precios se miden en tokens."),
    ("Ventana de contexto", "La cantidad máxima de tokens que un modelo puede tener en cuenta a la vez: instrucciones, "
                            "conversación, documentos y respuesta."),
]


def word_documents() -> None:
    for body, name, title in [
        (PRESS_RELEASE, "Nota de prensa – Trailhead 3 (borrador v2)", "Nota de prensa – Trailhead 3 (borrador v2)"),
        (OFFSITE_AGENDA, "Agenda offsite Lisboa", "Agenda – offsite de Experiencia de Cliente en Lisboa"),
        (KEYNOTE_SCRIPT, "Guion keynote – Data Summit Madrid", "Opening keynote – script (v3)"),
        (SECURITY_MINUTES, "Acta – comité de seguridad (16 oct)", "Acta – comité de seguridad, 16 de octubre"),
    ]:
        textutil(word(body), out(f"Documents/{name}.docx"), title)

    entries = "".join(f"<p><b>{esc(term)}</b><br>{esc(text_)}</p>" for term, text_ in AI_GLOSSARY)
    rich = f"""<!doctype html><html><head><meta charset="utf-8"><style>body{{font-family:'Helvetica Neue';
font-size:12pt}} h1{{font-size:22pt}} p{{margin:0 0 10pt}}</style></head><body><h1>Glosario de IA generativa</h1>
<p><i>Equipo de Experiencia de Cliente de Contoso · octubre de 2026</i></p>
<p>Las palabras que más se oyen en las reuniones del piloto del asistente, explicadas en una o dos frases y por orden
alfabético. Si falta alguna, díselo a Ana.</p>{entries}
<p><i>Fuentes: documentación pública de Microsoft Learn (Microsoft Foundry, Azure OpenAI y Microsoft Fabric) y
modelcontextprotocol.io. Contoso es una empresa ficticia.</i></p></body></html>"""
    textutil(rich, out("Documents/Glosario de IA generativa.rtf"), "Glosario de IA generativa")


# MARK: Slides


def heading(deck: Deck, title: str, kicker: str, notes: str):
    """A dark slide with the usual kicker, title, footer and number, in light colors: Deck.slide would draw them in
    the deck's ink, which is the background here."""
    p = deck.p
    s = deck.slide(title, notes=notes, dark=True)
    text(s, LEFT, 0.55, RIGHT - LEFT, 0.3, kicker.upper(), size=11, bold=True, color=p["accent"], tracking=150)
    place(s.shapes.title, LEFT, 0.84, RIGHT - LEFT, 0.8)
    write(s.shapes.title.text_frame, title, size=32, bold=True, color=p["text"], line=1.0)
    text(s, LEFT, 6.98, 9, 0.25, deck.footer, size=9.5, color=p["muted"])
    text(s, RIGHT - 1, 6.98, 1, 0.25, str(len(deck.titles)), size=9.5, color=p["muted"], align=PP_ALIGN.RIGHT)
    return s


def backdrop(slide, document: str, name: str) -> None:
    """A 1920×1080 page drawn full-bleed behind everything else on the slide."""
    path = WORK / f"{name}.png"
    save_png(render(document, (1920, 1080)), path)
    image = picture(slide, path, 0, 0, 13.333, 7.5)
    tree = slide.shapes._spTree
    tree.remove(image._element)
    tree.insert(2, image._element)


def tile(slide, svg: str, x: float, y: float, side: float, fill: str = "FFFFFF") -> None:
    """An official logo, whole and unaltered, centered on a square tile."""
    box(slide, x, y, side, side, fill)
    path, ratio = mark(svg)
    w = min(side * 0.62, side * 0.62 / ratio)
    picture(slide, path, x + (side - w) / 2, y + (side - w * ratio) / 2, w, w * ratio)


# The summit poster's halftone dots and violet glow, without its type.
SUMMIT_ART = """.dots{position:absolute;inset:0;
background:radial-gradient(circle,rgba(34,211,238,.5) 2px,transparent 2.4px) 0 0/20px 20px;
-webkit-mask-image:radial-gradient(1000px 820px at 100% 0%,#000 0%,transparent 72%)}
.beam{position:absolute;inset:0;
background:radial-gradient(900px 620px at 0% 100%,rgba(168,85,247,.36),transparent 70%)}"""


def keynote_deck() -> None:
    p = {"ink": "050510", "accent": "22D3EE", "soft": "0D1528", "line": "1E2A44", "muted": "7C86A8",
         "text": "F0FDFF", "light": "BAE6FD", "purple": "C084FC"}
    deck = Deck(p, "Data Summit Madrid · Opening keynote · Contoso Outdoor")
    notes = iter(said for *_, said in KEYNOTE_SCRIPT_SLIDES)
    width = (RIGHT - LEFT - 0.6) / 3

    s = deck.slide("How Contoso Outdoor handles returns with an AI agent", notes=next(notes), dark=True)
    backdrop(s, page('<div class="dots"></div><div class="beam"></div>', SUMMIT_ART, (1920, 1080), "#050510"),
             "summit-art")
    text(s, LEFT, 1.9, 10, 0.3, "DATA SUMMIT MADRID · OPENING KEYNOTE", size=12, bold=True, color=p["accent"],
         tracking=200)
    place(s.shapes.title, LEFT, 2.28, 11.2, 1.6)
    write(s.shapes.title.text_frame, "How Contoso Outdoor handles returns with an AI agent", size=46, bold=True,
          color=p["text"], line=1.0)
    text(s, LEFT, 4.05, 9.6, 0.5, "Lucía Romero · Head of Customer Experience, Contoso Outdoor", size=20,
         color=p["light"])
    box(s, LEFT, 4.75, 0.8, 0.055, fill=p["accent"])
    text(s, LEFT, 5.0, 9.6, 0.7, [("Thursday, November 12, 2026 · 9:15", {"bold": True, "color": p["text"]}),
                                  "Data Summit Madrid · 12 minutes"], size=15, color=p["muted"], after=3)

    s = heading(deck, "Eleven minutes, by hand", "Where we started", next(notes))
    text(s, LEFT - 0.08, 1.75, 4.6, 2.4, "11", size=170, bold=True, color=p["accent"], line=0.9)
    text(s, LEFT, 4.25, 4.4, 1.2, "minutes per return, every time, for almost the same answer.", size=20,
         color=p["light"], line=1.2)
    steps = [("Find the order", "In the back office, by name or by email."),
             ("Read the policy", "Damaged? Within 30 days? Who pays for the label?"),
             ("Open the return", "In the returns system, field by field."),
             ("Email a label", "And then answer the customer's next question.")]
    for index, (what, how) in enumerate(steps):
        top = 1.95 + index * 0.98
        box(s, 6.2, top + 0.12, 0.52, 0.52, p["soft"], p["line"], items=str(index + 1), size=16, bold=True,
            color=p["accent"], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        text(s, 7.0, top + 0.05, RIGHT - 7.0, 0.8, [(what, {"size": 19, "bold": True, "color": p["text"],
                                                            "after": 2}), how], size=14, color=p["muted"])
        if index < len(steps) - 1:
            box(s, 7.0, top + 0.93, RIGHT - 7.0, 0.012, fill=p["line"])
    box(s, LEFT, 5.95, RIGHT - LEFT, 0.72, p["soft"],
        items="September 2026 was our worst month: **2,914 returns**, online and in stores.", size=16,
        color=p["light"], anchor=MSO_ANCHOR.MIDDLE, pad=(0.35, 0.05, 0.35, 0.05))

    s = heading(deck, "Now 62 % resolve on their own", "Today", next(notes))
    stats = [("62 %", "resolve on their own", "The agent answers, opens the return and books the pickup."),
             ("4 min", "down from 11", "From the customer's first message to a booked pickup."),
             ("4.5 / 5", "customer satisfaction", "Customers rate every chat when it ends.")]
    for index, (number, what, detail) in enumerate(stats):
        box(s, LEFT + index * (width + 0.3), 1.95, width, 3.2, p["soft"], p["line"],
            items=[(number, {"size": 60, "bold": True, "color": p["accent"], "line": 1.0, "after": 10}),
                   (what, {"size": 19, "bold": True, "color": p["text"], "after": 8}), detail],
            size=14, color=p["muted"], line=1.22, pad=(0.35, 0.32, 0.35, 0.25))
    box(s, LEFT, 5.55, 0.06, 0.9, fill=p["purple"])
    text(s, LEFT + 0.35, 5.55, 11.0, 0.9, "The rest go to a person, with the whole case already written up: the order, "
                                          "the photos, the policy and what the customer wants.", size=18,
         color=p["light"], line=1.22, anchor=MSO_ANCHOR.MIDDLE)

    s = heading(deck, "Three things made it work", "How", next(notes))
    things = [("microsoft-fabric.svg", "Our data in one place",
               "Orders, returns and every policy together in OneLake, with **Microsoft Fabric**."),
              ("microsoft-foundry.svg", "An agent with tools",
               "Built with **Microsoft Foundry**: it looks up orders, opens returns and books pickups."),
              (None, "Evaluations", "**120 golden questions**, run before every release.")]
    for index, (svg, what, detail) in enumerate(things):
        x = LEFT + index * (width + 0.3)
        box(s, x, 1.95, width, 4.65, p["soft"], p["line"])
        if svg:
            tile(s, svg, x + 0.35, 2.3, 1.3)
        else:
            box(s, x + 0.35, 2.3, 1.3, 1.3, p["accent"], items="120", size=30, bold=True, color=p["ink"],
                align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        text(s, x + 0.35, 3.95, width - 0.7, 2.5, [
            (f"0{index + 1}", {"size": 12, "bold": True, "color": p["accent"], "tracking": 100, "after": 6}),
            (what, {"size": 21, "bold": True, "color": p["text"], "after": 8}), detail],
             size=15, color=p["light"], line=1.22)

    s = heading(deck, "Live demo: a tent with a broken pole", "Live demo", next(notes))
    box(s, LEFT, 1.95, RIGHT - LEFT, 1.0, p["soft"])
    box(s, LEFT, 1.95, 0.07, 1.0, fill=p["purple"])
    text(s, LEFT + 0.4, 1.95, 11.0, 1.0, [
        ("“My Trailhead 2 arrived with a broken pole. Order 48213.”", {"size": 21, "color": p["text"]}),
        ("Alex, a customer, in the chat on contoso.com", {"size": 12, "color": p["muted"], "before": 4})],
         anchor=MSO_ANCHOR.MIDDLE)
    calls = [("lookup_order", 'order_id="48213"', "Finds order #48213 and the tent in it."),
             ("file_search", '"damaged item"', "Return policy.pdf, p. 2: damaged items come back free."),
             ("create_return", 'reason="damaged"', "Opens return RMA-20931, with a prepaid label."),
             ("book_pickup", 'order_id="48213"', "Offers courier slots, Monday to Saturday.")]
    step = (RIGHT - LEFT - 1.5) / 4
    for index, (tool, argument, result) in enumerate(calls):
        x = LEFT + index * (step + 0.5)
        box(s, x, 3.35, step, 2.25, p["soft"], p["line"],
            items=[(tool, {"font": MONO, "size": 15, "bold": True, "color": p["accent"], "after": 4}),
                   (argument, {"font": MONO, "size": 10.5, "color": p["purple"], "after": 10}), result],
            size=13.5, color=p["light"], line=1.2, pad=(0.25, 0.25, 0.22, 0.2))
        if index < len(calls) - 1:
            link(s, x + step + 0.07, x + step + 0.43, 4.475, p["accent"], 1.75)
    text(s, LEFT, 5.95, RIGHT - LEFT, 0.6, "Watch the tool calls: each one is a step a person used to do by hand.",
         size=15, color=p["muted"])

    s = heading(deck, "What we learned", "Lessons", next(notes))
    lessons = [("Start with the boring questions", "Where's my order, how do I send it back: they're most of the "
                                                   "volume."),
               ("Write the evaluations first", "120 golden questions before the first prompt, run on every "
                                               "release."),
               ("Plan for rate limits", "On October 13 we got 429s for twenty minutes. Retries with backoff fixed "
                                        "it."),
               ("Hand over early", "As soon as a customer is upset, a person takes over, with the case written "
                                   "up.")]
    half = (RIGHT - LEFT - 0.3) / 2
    for index, (what, detail) in enumerate(lessons):
        row, column = divmod(index, 2)
        box(s, LEFT + column * (half + 0.3), 1.95 + row * 2.4, half, 2.15, p["soft"], p["line"],
            items=[(f"0{index + 1}", {"size": 12, "bold": True, "color": p["accent"], "tracking": 100, "after": 6}),
                   (what, {"size": 21, "bold": True, "color": p["text"], "after": 7}), detail],
            size=15, color=p["light"], line=1.22, pad=(0.38, 0.3, 0.38, 0.2))

    s = heading(deck, "What the data changed: Trailhead 3", "And then", next(notes))
    text(s, LEFT - 0.06, 1.85, 5.4, 2.0, "31 %", size=120, bold=True, color=p["accent"], line=0.9)
    text(s, LEFT, 3.75, 5.0, 1.3, "of September's returns had arrived damaged, many with a broken pole.", size=20,
         color=p["light"], line=1.22)
    text(s, LEFT, 5.2, 5.0, 0.6, "So the team redesigned the tent.", size=20, bold=True, color=p["text"])
    box(s, 6.6, 1.95, RIGHT - 6.6, 4.65, p["soft"], p["line"])
    text(s, 7.0, 2.25, RIGHT - 7.4, 0.7, "Trailhead 3", size=30, bold=True, color=p["text"])
    text(s, 7.0, 3.0, RIGHT - 7.4, 2.3, ["Sleeps three, and weighs **1.9 kg**",
                                         "Color-coded aluminum poles: up in about **three minutes**",
                                         "A **repair sleeve** in the bag, and five years of warranty on the poles",
                                         "**249\u00a0€**, in stores since November 3"],
         size=15, color=p["light"], line=1.2, after=8, bullet=("•", p["accent"], 0.25))
    for index, (color, name) in enumerate([("3D7358", "Forest green"), ("E2CFAD", "Sand")]):
        x = 7.0 + index * 2.2
        box(s, x, 5.65, 0.42, 0.42, color)
        text(s, x + 0.58, 5.65, 1.5, 0.42, name, size=13, color=p["light"], anchor=MSO_ANCHOR.MIDDLE)

    s = deck.slide("Thank you", notes=next(notes), dark=True)
    place(s.shapes.title, LEFT, 2.3, 10, 1.3)
    write(s.shapes.title.text_frame, "Thank you!", size=72, bold=True, color=p["text"], line=1.0)
    box(s, LEFT, 3.75, 0.8, 0.055, fill=p["accent"])
    text(s, LEFT, 4.0, 10, 1.2, [("Questions at the coffee break: I'll be next to the Fourth Coffee stand.",
                                  {"size": 21, "color": p["light"], "after": 8}),
                                 ("Lucía Romero · lucia.romero@contoso.com", {"size": 15, "color": p["muted"]})])
    text(s, LEFT, 6.98, 9, 0.25, deck.footer, size=9.5, color=p["muted"])

    deck.save(out("Slides/Data Summit Madrid – keynote.pptx"), "How Contoso Outdoor handles returns with an AI agent",
              dt.datetime(2026, 9, 30, 16, 5), dt.datetime(2026, 10, 13, 9, 40), 186, 14)


TENTS_CSS = """svg{position:absolute}.sand{left:380px;top:600px;width:600px;height:380px}
.forest{left:60px;top:720px;width:820px;height:519px}"""


def launch_deck() -> None:
    p = {"ink": "16261F", "accent": "2F5D46", "soft": "EEF3EE", "line": "D5DFD6", "muted": "5B6B62",
         "amber": "B45309", "amber_soft": "FDF3E4", "amber_line": "F0C98B"}
    deck = Deck(p, "Trailhead 3 · launch plan · Contoso Outdoor")
    width = (RIGHT - LEFT - 0.6) / 3

    s = deck.slide("Trailhead 3", notes="Fifteen minutes, three of us: Pablo for operations, Marta for the press, and "
                   "me. We agree on what happens between today and November 3, and on what we do about the poles.")
    box(s, 0, 0, 5.2, 7.5, fill=p["soft"])
    tents = (dome("#e2cfad", "#c9b089", "#a88d63", "#5b4a33", "sand").replace("<svg", '<svg class="sand"', 1)
             + dome("#3d7358", "#24493a", "#1d3a2e", "#16261f", "forest").replace("<svg", '<svg class="forest"', 1))
    save_png(render(page(tents, TENTS_CSS, (1040, 1500), "transparent"), (1040, 1500), transparent=True),
             WORK / "launch-tents.png")
    picture(s, WORK / "launch-tents.png", 0, 0, 5.2, 7.5)
    text(s, 6.0, 2.1, 6.6, 0.3, "CONTOSO OUTDOOR · LAUNCH PLAN", size=12, bold=True, color=p["accent"], tracking=200)
    place(s.shapes.title, 6.0, 2.45, 6.6, 0.95)
    write(s.shapes.title.text_frame, "Trailhead 3", size=54, bold=True, color=p["ink"], line=1.0)
    text(s, 6.0, 3.48, 6.0, 0.8, "From the poles in Zaragoza to every store, by Tuesday, November 3.", size=21,
         color=p["muted"], line=1.18)
    box(s, 6.0, 4.55, 0.8, 0.055, fill=p["accent"])
    text(s, 6.0, 4.8, 6.6, 0.7, [("Thursday, October 15, 2026", {"bold": True, "color": p["ink"]}),
                                 "Lucía, Pablo and Marta · Contoso Outdoor"], size=15, color=p["muted"], after=3)

    s = deck.slide("What we're launching", "The product", "Trailhead 3 in one slide. And remember legal's note: we "
                   "can't say it's the lightest tent on the market, because we can't prove it. We say it's 300 grams "
                   "lighter than Trailhead 2.")
    cards = [("1.9 kg, sleeps three", "300 g lighter than Trailhead 2, and 46 × 15 cm in its bag."),
             ("Up in three minutes", "Color-coded aluminum poles, two doors and two vestibules."),
             ("Made not to come back", "A repair sleeve in the bag and new packaging. Two years of warranty, five on "
                                       "the poles.")]
    for index, (what, detail) in enumerate(cards):
        box(s, LEFT + index * (width + 0.3), 1.95, width, 2.5, "FFFFFF", p["line"],
            items=[(f"0{index + 1}", {"size": 12, "bold": True, "color": p["accent"], "tracking": 100, "after": 7}),
                   (what, {"size": 20, "bold": True, "color": p["ink"], "after": 7}), detail],
            size=14.5, color=p["muted"], line=1.22, pad=(0.3, 0.28, 0.3, 0.2))
    box(s, LEFT, 4.75, 6.9, 1.85, p["soft"])
    text(s, LEFT + 0.4, 4.75, 2.6, 1.85, "249\u00a0€", size=46, bold=True, color=p["accent"], anchor=MSO_ANCHOR.MIDDLE)
    text(s, LEFT + 3.0, 5.0, 3.7, 0.8, "In stores and on contoso.com from **Tuesday, November 3**.", size=15,
         color=p["ink"], line=1.2)
    for index, (color, name) in enumerate([("3D7358", "Forest green"), ("E2CFAD", "Sand")]):
        x = LEFT + 3.0 + index * 1.75
        box(s, x, 5.9, 0.3, 0.3, color)
        text(s, x + 0.42, 5.9, 1.3, 0.3, name, size=12, color=p["muted"], anchor=MSO_ANCHOR.MIDDLE)
    box(s, LEFT + 7.2, 4.75, RIGHT - LEFT - 7.2, 1.85, p["amber_soft"], p["amber_line"],
        items=[("LEGAL", {"size": 11, "bold": True, "color": p["amber"], "tracking": 150, "after": 6}),
               "Don't write “the lightest tent on the market”: we can't prove it. Write “300 g lighter than "
               "Trailhead 2”."], size=14.5, color=p["ink"], line=1.22, pad=(0.3, 0.25, 0.3, 0.2))

    s = deck.slide("From the poles to the shelves", "Timeline", "Five dates. The one that worries me is the first: the "
                   "poles moved from October 13 to October 20. The product photos arrive on October 28.")
    milestones = [("Poles arrive", "Tue, Oct 20", "PO 5518, a week late: at the Zaragoza warehouse, 8:00–12:00. "
                                                   "Tracking FBK-RTM-772041."),
                  ("Assembly", "Oct 21–30", "Poles into the bags, with the repair sleeve: two shifts in Zaragoza."),
                  ("Stores stocked", "Fri, Oct 30", "The first tents in every store, boxed and under embargo."),
                  ("Launch", "Tue, Nov 3", "In stores and on contoso.com at 9:00 CET, with the press release."),
                  ("Keynote", "Thu, Nov 12", "Lucía opens Data Summit Madrid: why we redesigned the tent.")]
    step = (RIGHT - LEFT - 2.1) / 4
    box(s, LEFT + 0.1, 3.04, step * 4, 0.022, fill=p["line"])
    for index, (event, when, what) in enumerate(milestones):
        x = LEFT + index * step
        color = p["amber"] if index == 0 else p["accent"]
        text(s, x, 1.95, step - 0.2, 0.3, event.upper(), size=11, bold=True, color=color, tracking=120)
        text(s, x, 2.24, step - 0.2, 0.5, when, size=24, bold=True, color=p["ink"])
        dot(s, x + 0.1, 3.05, 0.2, p["accent"] if index == 3 else "FFFFFF", color, weight=2)
        text(s, x, 3.35, min(step - 0.25, RIGHT - x), 1.5, what, size=13.5, color=p["muted"], line=1.22)
    box(s, LEFT, 5.15, RIGHT - LEFT, 1.4, p["soft"])
    text(s, LEFT + 0.45, 5.15, 2.4, 1.4, "Oct 28", size=40, bold=True, color=p["accent"], anchor=MSO_ANCHOR.MIDDLE)
    text(s, LEFT + 2.9, 5.15, 8.4, 1.4, "The product photos arrive. Until then, the launch page and the press kit use "
                                         "the drawings from the spec sheet.", size=17, color=p["ink"], line=1.22,
         anchor=MSO_ANCHOR.MIDDLE)

    s = deck.slide("Risk: the poles arrive a week late", "Risk", "Tom Becker, from Fabrikam Logistics, called on "
                   "Tuesday: the container missed its vessel in Rotterdam. New delivery on Tuesday the 20th, between "
                   "eight and twelve, tracking FBK-RTM-772041. With two shifts in Zaragoza, Pablo says we make it. "
                   "The launch date stays November 3.")
    text(s, LEFT - 0.08, 1.75, 4.8, 2.4, "+7", size=170, bold=True, color=p["amber"], line=0.9)
    text(s, LEFT, 4.25, 4.6, 1.2, "days: the poles now arrive on Tuesday, October 20, not on the 13th.", size=20,
         color=p["ink"], line=1.2)
    facts = [("What happened", "PO 5518, the aluminum poles, missed its vessel in Rotterdam."),
             ("New delivery", "Tuesday, October 20, 8:00–12:00, at the Zaragoza warehouse."),
             ("Tracking", "FBK-RTM-772041 · Tom Becker, Fabrikam Logistics, +1 425 555 0142"),
             ("Impact", "Assembly gets 8 working days instead of 13: Pablo books a second shift.")]
    for index, (what, detail) in enumerate(facts):
        top = 1.95 + index * 0.98
        text(s, 6.2, top + 0.05, RIGHT - 6.2, 0.85, [(what, {"size": 17, "bold": True, "color": p["ink"], "after": 2}),
                                                     detail], size=14, color=p["muted"], line=1.18)
        if index < len(facts) - 1:
            box(s, 6.2, top + 0.93, RIGHT - 6.2, 0.012, fill=p["line"])
    box(s, LEFT, 5.95, RIGHT - LEFT, 0.72, p["amber_soft"],
        items="**The launch date doesn't move.** Plan B, if the poles slip again: contoso.com on November 3, the "
              "stores a week later.", size=15.5, color=p["ink"], anchor=MSO_ANCHOR.MIDDLE,
        pad=(0.35, 0.05, 0.35, 0.05))

    s = deck.slide("How people will hear about it", "Channels", "Five channels, one owner each. Only the teaser "
                   "goes out before Tuesday, November 3, at nine CET.")
    channels = [("contoso.com", "The launch page, counting down to November 3 at 9:00 CET.", "Sara"),
                ("Stores", "A pitched Trailhead 3 in every store, and pitching demos on Saturdays in November.",
                 "Pablo"),
                ("Press", "The press release, under embargo until November 3 at 9:00 CET.", "Marta"),
                ("Newsletter and social", "The teaser video on October 27, the setup guide on launch day.", "Marta"),
                ("Data Summit Madrid", "Lucía's opening keynote on November 12, with a tent on stage.", "Lucía")]
    for index, (channel, plan, owner) in enumerate(channels):
        top = 1.95 + index * 0.92
        text(s, LEFT, top, 3.4, 0.7, channel, size=19, bold=True, color=p["accent"], anchor=MSO_ANCHOR.MIDDLE)
        text(s, LEFT + 3.6, top, 6.5, 0.7, plan, size=15, color=p["ink"], line=1.18, anchor=MSO_ANCHOR.MIDDLE)
        text(s, RIGHT - 1.6, top, 1.6, 0.7, owner, size=13, bold=True, color=p["muted"], align=PP_ALIGN.RIGHT,
             anchor=MSO_ANCHOR.MIDDLE)
        if index < len(channels) - 1:
            box(s, LEFT, top + 0.81, RIGHT - LEFT, 0.012, fill=p["line"])

    s = deck.slide("Who does what", "Owners", "One name per line. If something is late, tell its owner first, and then "
                   "me.")
    owners = [("Pablo Serrano", "Operations · Zaragoza", "Poles in, tents assembled and every store stocked.",
               "October 30"),
              ("Marta Gil", "Press", "The press release, the embargo and the product photos.", "November 3"),
              ("Sara Navarro", "Software", "The launch page on contoso.com, and its countdown.", "November 2"),
              ("Ana Castillo", "Support", "The returns assistant knows all about Trailhead 3.", "October 30"),
              ("Javier Molina", "Data", "A launch dashboard in Power BI: sales and returns, day by day.", "November 3"),
              ("Lucía Romero", "Customer Experience", "The opening keynote at Data Summit Madrid.", "November 12")]
    for index, (name, team, task, due) in enumerate(owners):
        row, column = divmod(index, 3)
        box(s, LEFT + column * (width + 0.3), 1.95 + row * 2.4, width, 2.2, "FFFFFF", p["line"],
            items=[(name, {"size": 19, "bold": True, "color": p["ink"], "after": 1}),
                   (team, {"size": 12, "color": p["muted"], "after": 9}), task,
                   (f"Due {due}", {"size": 12, "bold": True, "color": p["accent"], "before": 9})],
            size=14, color=p["ink"], line=1.2, pad=(0.3, 0.26, 0.3, 0.2))

    s = deck.slide("Next steps", "This week", "Five things for this week. I'll send these notes after the meeting.")
    todo = [("Confirm Tuesday's delivery with Fabrikam, and who signs in Zaragoza", "Pablo", "Fri, Oct 16"),
            ("Legal review of the press release, v2", "Marta", "Fri, Oct 16"),
            ("Decide: the assistant's quote in the press release, or in the keynote?", "Lucía", "Fri, Oct 16"),
            ("Book a second assembly shift from October 21", "Pablo", "Mon, Oct 19"),
            ("Add the spec sheet and the setup guide to the returns assistant", "Ana", "Fri, Oct 23")]
    for index, (what, who, when) in enumerate(todo):
        top = 1.95 + index * 0.9
        box(s, LEFT, top + 0.05, 0.34, 0.34, None, p["accent"], 1.75)
        text(s, LEFT + 0.6, top, 8.9, 0.45, what, size=16, color=p["ink"])
        text(s, 10.35, top, 0.9, 0.45, who, size=14, bold=True, color=p["muted"])
        text(s, RIGHT - 1.3, top, 1.3, 0.45, when, size=14, bold=True, color=p["accent"], align=PP_ALIGN.RIGHT)
        if index < len(todo) - 1:
            box(s, LEFT, top + 0.68, RIGHT - LEFT, 0.012, fill=p["line"])
    box(s, LEFT, 6.2, RIGHT - LEFT, 0.55, p["soft"],
        items="Next check-in: **Monday, October 19, at 10:00**. Lucía is in Lisbon from Wednesday 21 to Friday 23.",
        size=14, color=p["ink"], anchor=MSO_ANCHOR.MIDDLE, pad=(0.35, 0, 0.35, 0))

    deck.save(out("Slides/Trailhead 3 – launch plan.pptx"), "Trailhead 3 – launch plan",
              dt.datetime(2026, 10, 5, 10, 20), dt.datetime(2026, 10, 15, 14, 5), 132, 9)


# A wall of blue-and-white tiles, as in Lisbon: rings where four tiles meet, a flower in the middle of each.
AZULEJO = """<svg class="wall" viewBox="0 0 900 1080"><defs><pattern id="az" width="180" height="180"
patternUnits="userSpaceOnUse"><rect width="180" height="180" fill="#f3f1ea"/>
<g fill="#1e3a8a"><circle r="50"/><circle cx="180" r="50"/><circle cy="180" r="50"/><circle cx="180" cy="180" r="50"/>
<circle cx="90" r="13"/><circle cx="90" cy="180" r="13"/><circle cy="90" r="13"/><circle cx="180" cy="90" r="13"/></g>
<g fill="#f3f1ea"><circle r="36"/><circle cx="180" r="36"/><circle cy="180" r="36"/>
<circle cx="180" cy="180" r="36"/></g>
<g fill="#2f5fd0"><circle r="16"/><circle cx="180" r="16"/><circle cy="180" r="16"/>
<circle cx="180" cy="180" r="16"/></g>
<g fill="#1d4ed8" transform="translate(90 90)"><ellipse cy="-38" rx="14" ry="30"/><ellipse cy="38" rx="14" ry="30"/>
<ellipse cx="-38" rx="30" ry="14"/><ellipse cx="38" rx="30" ry="14"/></g>
<g fill="#93a5dc" transform="translate(90 90) rotate(45)"><ellipse cy="-34" rx="8" ry="18"/>
<ellipse cy="34" rx="8" ry="18"/><ellipse cx="-34" rx="18" ry="8"/><ellipse cx="34" rx="18" ry="8"/></g>
<circle cx="90" cy="90" r="14" fill="#f2b705"/><circle cx="90" cy="90" r="5" fill="#1e3a8a"/>
<path d="M0 0H180V180" fill="none" stroke="#d5d9e4" stroke-width="2"/></pattern></defs>
<rect width="900" height="1080" fill="url(#az)"/></svg><div class="glaze"></div><div class="edge"></div>"""
AZULEJO_CSS = """.wall,.glaze{position:absolute;right:0;top:0;width:900px;height:1080px}
.glaze{background:radial-gradient(700px 520px at 80% 8%,rgba(255,255,255,.2),transparent 70%)}
.edge{position:absolute;left:1020px;top:0;width:90px;height:1080px;
background:linear-gradient(90deg,rgba(10,18,40,.6),rgba(10,18,40,0))}"""


def offsite_deck() -> None:
    p = {"ink": "14213D", "accent": "1D4ED8", "soft": "EEF2FB", "line": "D6DEF0", "muted": "5B6478",
         "text": "F5F7FF", "light": "C7D2FE", "yellow": "F2B705", "yellow_soft": "FEF6D8", "yellow_line": "F2D27A"}
    deck = Deck(p, "Offsite Lisboa · 21–23 de octubre · Experiencia de Cliente, Contoso")
    half = (RIGHT - LEFT - 0.3) / 2
    width = (RIGHT - LEFT - 0.6) / 3

    s = deck.slide("Offsite Lisboa", notes="Bienvenida. Tres días en Lisboa, del miércoles 21 al viernes 23. Aquí está "
                   "todo: vuelos, hotel, agenda y lo práctico.", dark=True)
    backdrop(s, page(AZULEJO, AZULEJO_CSS, (1920, 1080), "#14213d"), "offsite-azulejos")
    text(s, LEFT, 2.0, 6.0, 0.3, "EXPERIENCIA DE CLIENTE · OFFSITE", size=12, bold=True, color=p["yellow"],
         tracking=200)
    place(s.shapes.title, LEFT, 2.38, 6.0, 0.95)
    write(s.shapes.title.text_frame, "Offsite Lisboa", size=54, bold=True, color=p["text"], line=1.0)
    text(s, LEFT, 3.45, 5.9, 0.8, "Del miércoles 21 al viernes 23 de octubre de 2026. Catorce personas, tres días.",
         size=21, color=p["light"], line=1.18)
    box(s, LEFT, 4.55, 0.8, 0.055, fill=p["yellow"])
    text(s, LEFT, 4.8, 5.9, 0.7, [("Vuelos, hotel y agenda", {"bold": True, "color": p["text"]}),
                                  "Organiza Ana Castillo · actualizada el 15 de octubre"], size=15, color=p["light"],
         after=3)

    s = deck.slide("El viaje", "Vuelos y hotel", "Volamos todos juntos: el BY 471 el miércoles a las 08:35 desde la "
                   "T4, y la vuelta el viernes en el BY 476, a las 19:40. La reserva es MT-4Q7ZK.")
    flights = [("IDA · MIÉRCOLES 21 DE OCTUBRE", "MAD → LIS", "BY 471 · Blue Yonder Airlines",
                "Sale a las 08:35 de Madrid T4 y llega a las 08:55 a Lisboa T1, hora local."),
               ("VUELTA · VIERNES 23 DE OCTUBRE", "LIS → MAD", "BY 476 · Blue Yonder Airlines",
                "Sale a las 19:40 de Lisboa T1 y llega a las 21:55 a Madrid T4.")]
    for index, (when, route, flight, detail) in enumerate(flights):
        x = LEFT + index * (half + 0.3)
        box(s, x, 1.95, half, 2.45, "FFFFFF", p["line"])
        box(s, x, 1.95, 0.08, 2.45, fill=p["yellow"] if index == 0 else p["accent"])
        text(s, x + 0.45, 2.2, half - 0.8, 2.1, [
            (when, {"size": 11, "bold": True, "color": p["accent"], "tracking": 120, "after": 4}),
            (route, {"size": 40, "bold": True, "color": p["ink"], "line": 1.0, "after": 6}),
            (flight, {"size": 15, "bold": True, "color": p["ink"], "after": 3}), detail],
             size=14, color=p["muted"], line=1.2)
    facts = [("MT-4Q7ZK", "La reserva de todo el equipo, con Blue Yonder Airlines."),
             ("–1 h", "Lisboa va una hora por detrás de Madrid. Las horas de Lisboa son locales."),
             ("≤ 140\u00a0€", "Por noche: un hotel en la Baixa, con desayuno y una sala para las sesiones.")]
    for index, (value, detail) in enumerate(facts):
        box(s, LEFT + index * (width + 0.3), 4.7, width, 1.9, p["soft"],
            items=[(value, {"size": 28, "bold": True, "color": p["accent"], "after": 6}), detail],
            size=14, color=p["ink"], line=1.22, pad=(0.3, 0.28, 0.3, 0.2))

    def day(title: str, kicker: str, notes: str, rows: list[tuple[str, str, str]], callout=None) -> None:
        s = deck.slide(title, kicker, notes)
        for index, (when, what, who) in enumerate(rows):
            top = 1.95 + index * 0.78
            text(s, LEFT, top, 2.2, 0.55, when, size=20, bold=True, color=p["accent"], anchor=MSO_ANCHOR.MIDDLE)
            text(s, LEFT + 2.4, top, 7.0, 0.55, what, size=18, color=p["ink"], anchor=MSO_ANCHOR.MIDDLE)
            text(s, RIGHT - 2.4, top, 2.4, 0.55, who, size=14, color=p["muted"], align=PP_ALIGN.RIGHT,
                 anchor=MSO_ANCHOR.MIDDLE)
            if index < len(rows) - 1:
                box(s, LEFT, top + 0.66, RIGHT - LEFT, 0.012, fill=p["line"])
        if callout:
            label, body, warm = callout
            box(s, LEFT, 5.15, RIGHT - LEFT, 1.45, p["yellow_soft"] if warm else p["soft"],
                p["yellow_line"] if warm else None,
                items=[(label, {"size": 11, "bold": True, "color": "8A6100" if warm else p["accent"],
                                 "tracking": 150, "after": 6}), *([body] if isinstance(body, str) else body)],
                size=16, color=p["ink"], line=1.22, anchor=MSO_ANCHOR.MIDDLE, pad=(0.4, 0.15, 0.4, 0.15))

    day("Miércoles 21", "Día 1", "El miércoles empezamos a las once, con los resultados del piloto, y por la tarde "
        "hacemos el taller sobre el asistente en 2027.",
        [("08:35", "Vuelo BY 471, desde Madrid T4", "Llega a las 08:55"),
         ("10:30", "Llegada al hotel y café", ""),
         ("11:00–13:00", "Bienvenida y resultados del piloto", "Lucía"),
         ("13:00", "Comida", ""),
         ("14:30–17:30", "Taller: el asistente en 2027", "Ana y Sara"),
         ("Tarde", "Libre", "")])
    day("Jueves 22", "Día 2", "El jueves por la mañana, estrategia 2027. Por la tarde visitamos la futura oficina. Y a "
        "las ocho y media, cena en Alfama.",
        [("9:00–13:00", "Estrategia 2027", "Lucía"),
         ("13:00", "Comida", ""),
         ("14:30–17:00", "Visita a la futura oficina de Lisboa", "Lucía y Pablo"),
         ("20:30", "Cena en Alfama, junto al mirador", "Todos")],
        ("CENA DEL JUEVES · 20:30", ["Somos 14: 3 menús vegetarianos y 1 sin gluten.",
                                     "¿Alguna otra alergia o intolerancia? Díselo a Ana **antes del martes 20 a las "
                                     "12:00**."], True))
    day("Viernes 23", "Día 3", "El viernes cerramos con el plan del primer trimestre. Las maletas se quedan en el "
        "hotel, y a las cinco y media salen los taxis al aeropuerto.",
        [("9:30–12:30", "Plan del primer trimestre y retro del offsite", "Ana"),
         ("12:30", "Check-out: las maletas se quedan en el hotel", ""),
         ("17:30", "Taxi al aeropuerto, desde el hotel", ""),
         ("19:40", "Vuelo BY 476 a Madrid", "Llega a las 21:55")],
        ("VUELTA · BY 476", "Sale a las 19:40 de Lisboa T1 y llega a las 21:55 a Madrid T4. El taxi al aeropuerto "
                            "entra en los gastos: guarda el recibo.", False))

    s = deck.slide("Información práctica", "A tener en cuenta", "Lo práctico: los gastos, según la FIN-04; calzado "
                   "cómodo; y recordad que en Lisboa es una hora menos.")
    tips = [("Gastos", "Según la política de viajes FIN-04: comidas hasta 45\u00a0€ al día. El taxi al aeropuerto "
                       "entra."),
            ("Calzado cómodo", "Del hotel al restaurante hay diez minutos andando, cuesta arriba."),
            ("Hora local", "Lisboa va una hora por detrás de Madrid. Todas las horas de esta agenda son locales."),
            ("La oficina de Lisboa", "Abre en marzo de 2027. Esperamos ser 40 personas allí a finales de 2027.")]
    for index, (what, detail) in enumerate(tips):
        row, column = divmod(index, 2)
        box(s, LEFT + column * (half + 0.3), 1.95 + row * 2.4, half, 2.15, "FFFFFF", p["line"],
            items=[(f"0{index + 1}", {"size": 12, "bold": True, "color": p["accent"], "tracking": 100, "after": 6}),
                   (what, {"size": 21, "bold": True, "color": p["ink"], "after": 7}), detail],
            size=15, color=p["muted"], line=1.22, pad=(0.38, 0.3, 0.38, 0.2))

    deck.save(out("Slides/Offsite Lisboa – agenda.pptx"), "Offsite Lisboa – agenda", dt.datetime(2026, 10, 13, 15, 30),
              dt.datetime(2026, 10, 15, 11, 10), 64, 6)


def copilot_deck() -> None:
    p = {"ink": "1F2328", "accent": "8250DF", "soft": "F6F8FA", "line": "D0D7DE", "muted": "57606A"}
    deck = Deck(p, "GitHub Copilot rollout · Contoso engineering")
    width = (RIGHT - LEFT - 0.6) / 3

    s = deck.slide("GitHub Copilot rollout", notes="This is the plan to give GitHub Copilot to all 420 engineers by "
                   "December 1, in three waves, with the guardrails the security committee asked for.")
    box(s, 0, 0, 5.2, 7.5, fill=p["soft"])
    path, ratio = mark("github-copilot-lockup.svg")
    w = min(3.8, 3.0 / ratio)
    picture(s, path, (5.2 - w) / 2, (7.5 - w * ratio) / 2, w, w * ratio)
    text(s, 6.0, 1.75, 6.6, 0.3, "CONTOSO ENGINEERING · ROLLOUT PLAN", size=12, bold=True, color=p["accent"],
         tracking=200)
    place(s.shapes.title, 6.0, 2.1, 6.6, 1.6)
    write(s.shapes.title.text_frame, "GitHub Copilot rollout", size=48, bold=True, color=p["ink"], line=1.0)
    text(s, 6.0, 4.05, 6.0, 0.8, "420 engineers in three waves, from October 19 to December 1.", size=21,
         color=p["muted"], line=1.18)
    box(s, 6.0, 5.05, 0.8, 0.055, fill=p["accent"])
    text(s, 6.0, 5.3, 6.6, 0.7, [("Tuesday, October 13, 2026", {"bold": True, "color": p["ink"]}),
                                 "Sara Navarro · Software engineering"], size=15, color=p["muted"], after=3)

    s = deck.slide("Three waves, 420 engineers", "The plan", "Three waves. Sixty engineers on October 19, the teams in "
                   "the pilot; 160 more on November 2; the last 200 on November 16. Everyone has it by December 1, and "
                   "the developer survey goes out on December 15.")
    waves = [("WAVE 1 · MON, OCT 19", 60, "The pilot teams: the returns assistant, the web and data.", "5A32A3"),
             ("WAVE 2 · MON, NOV 2", 160, "Every product team, each starting with a 90-minute kickoff.", "8250DF"),
             ("WAVE 3 · MON, NOV 16", 200, "Everyone else, contractors with a Contoso account included.", "C4A9F4")]
    x, total = LEFT, 0
    for index, (when, count, who, color) in enumerate(waves):
        box(s, LEFT + index * (width + 0.3), 1.95, width, 3.0, "FFFFFF", p["line"],
            items=[(when, {"size": 11, "bold": True, "color": p["accent"], "tracking": 120, "after": 8}),
                   (str(count), {"size": 60, "bold": True, "color": p["ink"], "line": 1.0, "after": 2}),
                   ("engineers", {"size": 15, "bold": True, "color": p["muted"], "after": 10}), who],
            size=14, color=p["ink"], line=1.22, pad=(0.32, 0.3, 0.32, 0.2))
        span = (RIGHT - LEFT) * count / 420
        total += count
        box(s, x, 5.3, span - 0.04, 0.2, color)
        text(s, x, 5.56, span - 0.04, 0.3, str(total), size=12, bold=True, color=p["muted"], align=PP_ALIGN.RIGHT)
        x += span
    text(s, LEFT, 6.0, RIGHT - LEFT, 0.6, "**420 engineers** with GitHub Copilot by **December 1**. The developer "
                                          "survey goes out on **December 15**.", size=17, color=p["ink"])

    s = deck.slide("What every engineer gets", "Features", "Six things, all in the product today. In the kickoff we "
                   "show agent mode and the coding agent live, because they change how a team works the most.")
    features = [("Code completions", "Suggestions as you type, in VS Code, Visual Studio, JetBrains IDEs, Xcode and "
                                     "more."),
                ("Copilot Chat", "Ask about your code in the editor or on GitHub.com: explain it, fix it, test it."),
                ("Agent mode", "Describe a change: Copilot edits the files, runs commands and iterates until it's "
                               "done."),
                ("Copilot code review", "Add Copilot as a reviewer on a pull request, and it comments with suggested "
                                        "changes."),
                ("Copilot coding agent", "Assign it an issue: it works in the background and opens a pull request."),
                ("Custom instructions", "A .github/copilot-instructions.md file tells Copilot our style and "
                                        "libraries.")]
    for index, (name, detail) in enumerate(features):
        row, column = divmod(index, 3)
        box(s, LEFT + column * (width + 0.3), 1.95 + row * 2.2, width, 1.95, "FFFFFF", p["line"],
            items=[(f"0{index + 1}", {"size": 12, "bold": True, "color": p["accent"], "tracking": 100, "after": 5}),
                   (name, {"size": 18, "bold": True, "color": p["ink"], "after": 5}), detail],
            size=13.5, color=p["muted"], line=1.2, pad=(0.3, 0.25, 0.3, 0.2))
    text(s, LEFT, 6.4, RIGHT - LEFT, 0.3, "Source: GitHub Copilot documentation, docs.github.com.", size=11,
         color=p["muted"])

    s = deck.slide("Guardrails, on before wave 1", "Security", "Three policies, set in the organization before anyone "
                   "gets a license. Marc owns them, and they're on the security committee's agenda on Friday.")
    rules = [("Block public code matches", "Suggestions that match public code are blocked, for the whole "
                                           "organization."),
             ("Content exclusion", "Configuration files with secrets stay out of Copilot's context."),
             ("A person approves every merge", "Copilot can write a pull request and review it, but a person "
                                               "approves it.")]
    for index, (rule, detail) in enumerate(rules):
        x = LEFT + index * (width + 0.3)
        box(s, x, 1.95, width, 3.3, "FFFFFF", p["line"])
        dot(s, x + 0.62, 2.62, 0.64, p["accent"], items=str(index + 1), size=20, bold=True, color="FFFFFF",
            align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        text(s, x + 0.3, 3.2, width - 0.6, 2.0, [(rule, {"size": 19, "bold": True, "color": p["ink"], "after": 8}),
                                                  detail], size=14.5, color=p["muted"], line=1.22)
    box(s, LEFT, 5.6, RIGHT - LEFT, 0.95, p["soft"],
        items="Owner: **Marc Puig**, before wave 1 on **October 19**. On the security committee's agenda for Friday, "
              "October 16.", size=16, color=p["ink"], anchor=MSO_ANCHOR.MIDDLE, pad=(0.4, 0, 0.4, 0))

    s = deck.slide("How we'll know it works", "Metrics", "Three numbers. If satisfaction comes in under three and a "
                   "half, we stop the rollout and ask why before going on.")
    stats = [("80 %", "weekly active users", "Of the 420 licenses, by December 15, from the Copilot metrics API."),
             ("–20 %", "pull request cycle time", "The median from opening a pull request to merging it, against "
                                                  "September."),
             ("4 / 5", "developer satisfaction", "In the December 15 survey. Under 3.5, we stop and ask why.")]
    for index, (number, what, detail) in enumerate(stats):
        box(s, LEFT + index * (width + 0.3), 1.95, width, 3.2, p["soft"], p["line"],
            items=[(number, {"size": 56, "bold": True, "color": p["accent"], "line": 1.0, "after": 10}),
                   (what, {"size": 19, "bold": True, "color": p["ink"], "after": 8}), detail],
            size=14, color=p["muted"], line=1.22, pad=(0.35, 0.32, 0.35, 0.25))
    box(s, LEFT, 5.55, 0.06, 0.9, fill=p["accent"])
    text(s, LEFT + 0.35, 5.55, 11.0, 0.9, "The baseline is September 2026, from our pull request data in Microsoft "
                                          "Fabric. Javier builds the dashboard.", size=17, color=p["ink"], line=1.22,
         anchor=MSO_ANCHOR.MIDDLE)

    s = deck.slide("Training and support", "Enablement", "Nobody gets a license without the kickoff. After that, "
                   "office hours every Thursday and a champion in every team.")
    support = [("Kickoff", "90 minutes on the Monday each wave starts: setup, chat, agent mode and the guardrails.",
                "Sara"),
               ("Office hours", "Every Thursday at 16:00, on Teams, until December 17.", "Champions"),
               ("Champions", "One per team from wave 1: they answer questions and collect feedback.", "Team leads"),
               ("Learning paths", "GitHub Skills and the GitHub Copilot courses on Microsoft Learn, self-paced.", ""),
               ("Cheat sheets", "“Chuleta de prompts para Copilot” and the video “GitHub Copilot – 5 tips”.", "")]
    for index, (what, plan, owner) in enumerate(support):
        top = 1.95 + index * 0.92
        text(s, LEFT, top, 3.0, 0.7, what, size=19, bold=True, color=p["accent"], anchor=MSO_ANCHOR.MIDDLE)
        text(s, LEFT + 3.2, top, 7.0, 0.7, plan, size=15, color=p["ink"], line=1.18, anchor=MSO_ANCHOR.MIDDLE)
        text(s, RIGHT - 1.6, top, 1.6, 0.7, owner, size=13, bold=True, color=p["muted"], align=PP_ALIGN.RIGHT,
             anchor=MSO_ANCHOR.MIDDLE)
        if index < len(support) - 1:
            box(s, LEFT, top + 0.81, RIGHT - LEFT, 0.012, fill=p["line"])

    s = deck.slide("Before October 19", "Next steps", "Five things before wave 1. The policies come first: no "
                   "policies, no licenses.")
    todo = [("Turn on the policies: public code matches, content exclusion, reviews", "Marc", "Fri, Oct 16"),
            ("Assign the 60 licenses of wave 1", "Sara", "Fri, Oct 16"),
            ("Book the wave 1 kickoff: Monday, October 19, at 10:00", "Sara", "Wed, Oct 14"),
            ("Name a champion in each pilot team", "Leads", "Fri, Oct 16"),
            ("A September baseline for pull request cycle time, in Fabric", "Javier", "Fri, Oct 23")]
    for index, (what, who, when) in enumerate(todo):
        top = 1.95 + index * 0.9
        box(s, LEFT, top + 0.05, 0.34, 0.34, None, p["accent"], 1.75)
        text(s, LEFT + 0.6, top, 8.9, 0.45, what, size=16, color=p["ink"])
        text(s, 10.35, top, 0.9, 0.45, who, size=14, bold=True, color=p["muted"])
        text(s, RIGHT - 1.3, top, 1.3, 0.45, when, size=14, bold=True, color=p["accent"], align=PP_ALIGN.RIGHT)
        if index < len(todo) - 1:
            box(s, LEFT, top + 0.68, RIGHT - LEFT, 0.012, fill=p["line"])
    box(s, LEFT, 6.2, RIGHT - LEFT, 0.55, p["soft"], items="Questions: **#copilot-rollout** on Teams, or ask Sara.",
        size=14, color=p["ink"], anchor=MSO_ANCHOR.MIDDLE, pad=(0.35, 0, 0.35, 0))

    deck.save(out("Slides/GitHub Copilot rollout – engineering.pptx"), "GitHub Copilot rollout – engineering",
              dt.datetime(2026, 10, 8, 12, 0), dt.datetime(2026, 10, 12, 15, 20), 97, 11)


@builder
def more_documents() -> None:
    """Documents/: the PDFs the returns assistant cites, policies, cheat sheets, Word documents and a glossary."""
    return_policy()
    shipping_options()
    spec_sheet()
    travel_policy()
    summit_agenda()
    kql_cheat_sheet()
    prompt_cheat_sheet()
    word_documents()


@builder
def more_slides() -> None:
    """Slides/: the Data Summit keynote, Trailhead 3's launch plan, the Lisbon offsite and GitHub Copilot's rollout."""
    keynote_deck()
    launch_deck()
    offsite_deck()
    copilot_deck()
